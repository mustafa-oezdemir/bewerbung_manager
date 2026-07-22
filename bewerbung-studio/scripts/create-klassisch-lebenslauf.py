#!/usr/bin/env python3
"""Create the bundled Klassisch visual/ATS Word templates and card preview."""

from __future__ import annotations

import argparse
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

A4_WIDTH_MM = 210
A4_HEIGHT_MM = 297
FONT = "Arial"
PRIMARY = "2B2F32"
ACCENT = "00AFC5"
HEADING = "5A6267"
TEXT = "3F484D"
MUTED = "68747A"
SOFT = "CDEFF3"
BORDER = "D5DBDE"
ATS_BLUE = "173B63"
CONTENT_WIDTH_DXA = 10_224

# compact_reference_guide resolved with the named klassisch_cv_a4_override:
# A4, 16/15/17/15 mm margins, Arial 8.5 pt, 1.15 line spacing,
# fixed-DXA borderless layout tables, and conservative gray/cyan hierarchy.


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size: float, color: str = TEXT, bold: bool = False) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:cs"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def configure_paragraph(
    paragraph,
    *,
    before: float = 0,
    after: float = 0,
    line: float = 1.15,
    keep_next: bool = False,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def paragraph_for(parent):
    if hasattr(parent, "paragraphs") and len(parent.paragraphs) == 1:
        paragraph = parent.paragraphs[0]
        if not paragraph.text and not paragraph.runs:
            return paragraph
    return parent.add_paragraph()


def add_text(
    parent,
    text: str,
    *,
    size: float = 8.5,
    color: str = TEXT,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.15,
    keep_next: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    paragraph = paragraph_for(parent)
    paragraph.alignment = align
    configure_paragraph(
        paragraph,
        before=before,
        after=after,
        line=line,
        keep_next=keep_next,
    )
    set_run(paragraph.add_run(text), size, color, bold)
    return paragraph


def add_hyperlink(parent, text: str, url: str, *, size=8.0, color=TEXT):
    paragraph = paragraph_for(parent)
    configure_paragraph(paragraph, after=0, line=1.05)
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.extend([fonts, color_node, size_node, size_cs])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return paragraph


def set_cell_margins(cell, top=50, start=40, bottom=50, end=40) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 0) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(A4_WIDTH_MM)
    section.page_height = Mm(A4_HEIGHT_MM)
    section.top_margin = Mm(16)
    section.right_margin = Mm(15)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(15)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(5)


def configure_styles(doc: Document, primary: str, body_size=8.5) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = rgb(TEXT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    bullet = doc.styles["List Bullet"]
    bullet.font.name = FONT
    bullet._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    bullet.font.size = Pt(body_size)
    bullet.font.color.rgb = rgb(TEXT)
    bullet.paragraph_format.left_indent = Mm(4.3)
    bullet.paragraph_format.first_line_indent = Mm(-2.2)
    bullet.paragraph_format.space_after = Pt(0.4)
    bullet.paragraph_format.line_spacing = 1.1
    for style_name, size in (("Heading 1", 10.4), ("Heading 2", 9.8), ("Heading 3", 9.0)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(primary)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3)


def add_section_title(parent, title: str, *, ats=False) -> None:
    paragraph = add_text(
        parent,
        title,
        size=10.5 if ats else 10.4,
        color=ATS_BLUE if ats else HEADING,
        bold=True,
        before=8 if ats else 7,
        after=5 if ats else 5,
        line=1.0,
        keep_next=True,
    )
    if ats:
        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "5")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "AEB8BF")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)


def add_real_bullet(parent, text: str, *, size=8.2, color=TEXT) -> None:
    paragraph = parent.add_paragraph(style="List Bullet")
    configure_paragraph(paragraph, after=0.4, line=1.1)
    paragraph.paragraph_format.keep_together = True
    set_run(paragraph.add_run(text), size, color)


def cubic_points(p0, p1, p2, p3, count=64):
    points = []
    for index in range(count + 1):
        t = index / count
        u = 1 - t
        x = u**3 * p0[0] + 3 * u**2 * t * p1[0] + 3 * u * t**2 * p2[0] + t**3 * p3[0]
        y = u**3 * p0[1] + 3 * u**2 * t * p1[1] + 3 * u * t**2 * p2[1] + t**3 * p3[1]
        points.append((int(x), int(y)))
    return points


def draw_wave_background(path: Path) -> None:
    image = Image.new("RGBA", (1240, 1754), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    fill = (205, 239, 243, 255)
    white = (255, 255, 255, 235)
    top = cubic_points((220, 0), (390, 220), (760, 40), (1240, 300))
    draw.polygon([(220, 0), (1240, 0), (1240, 300), *reversed(top)], fill=fill)
    for offset in (0, 34):
        line = cubic_points((420 - offset, -20), (520, 120 + offset), (780, 15 + offset), (1240, 150 + offset))
        draw.line(line, fill=white, width=2)
    draw.ellipse((1100, -80, 1240, 75), outline=white, width=2)
    bottom = cubic_points((0, 1480), (230, 1460), (330, 1680), (520, 1754))
    draw.polygon([(0, 1480), *bottom, (0, 1754)], fill=fill)
    for offset in (0, 35):
        line = cubic_points((-20, 1570 + offset), (150, 1505 + offset), (310, 1690 + offset), (455, 1765 + offset))
        draw.line(line, fill=white, width=2)
    draw.ellipse((-45, 1660, 110, 1815), outline=white, width=2)
    image.save(path)


def create_placeholder_photo(path: Path) -> None:
    image = Image.new("RGBA", (420, 420), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse((0, 0, 419, 419), fill="#DCE5E7")
    draw.ellipse((110, 58, 310, 258), fill="#68747A")
    draw.ellipse((60, 225, 360, 525), fill="#68747A")
    image.save(path)


def set_photo_shape(inline_shape) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", "PROFILFOTO")
    doc_pr.set("title", "PROFILFOTO")
    geometry = inline_shape._inline.find(".//" + qn("a:prstGeom"))
    if geometry is not None:
        geometry.set("prst", "ellipse")


def add_header_background(doc: Document, image_path: Path) -> None:
    header = doc.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(
        str(image_path), width=Mm(210), height=Mm(297)
    )
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", "0"),
        ("behindDoc", "1"),
        ("locked", "0"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
    ):
        anchor.set(key, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for tag in ("positionH", "positionV"):
        position = OxmlElement(f"wp:{tag}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    for element in (inline.find(qn("wp:extent")), inline.find(qn("wp:effectExtent"))):
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element in (
        inline.find(qn("wp:docPr")),
        inline.find(qn("wp:cNvGraphicFramePr")),
        inline.find(qn("a:graphic")),
    ):
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def add_page_number(paragraph) -> None:
    for node_name, attributes, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "1"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        node = OxmlElement(node_name)
        for key, value in attributes.items():
            node.set(qn(key), value)
        if text is not None:
            node.text = text
        paragraph.add_run()._r.append(node)


def add_footer(doc: Document, *, ats=False) -> None:
    footer = doc.sections[0].footer
    table = footer.add_table(rows=1, cols=2, width=Mm(180))
    set_table_geometry(table, [5_112, 5_112])
    remove_table_borders(table)
    left, right = table.rows[0].cells
    add_hyperlink(left, "{{WEBSITE}}", "{{WEBSITE}}", size=7.0, color=MUTED)
    paragraph = add_text(
        right,
        "Seite ",
        size=7.0,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    add_page_number(paragraph)
    if ats:
        footer.is_linked_to_previous = False


def add_contact_grid(parent) -> None:
    table = parent.add_table(rows=2, cols=3)
    set_table_geometry(table, [3_408, 3_408, 3_408])
    remove_table_borders(table)
    cells = [cell for row in table.rows for cell in row.cells]
    values = [
        ("{{HEADER_KONTAKT_1}}", "tel:{{HEADER_KONTAKT_1}}"),
        ("{{HEADER_KONTAKT_2}}", "mailto:{{HEADER_KONTAKT_2}}"),
        ("{{HEADER_KONTAKT_3}}", "{{HEADER_KONTAKT_3}}"),
        ("{{HEADER_KONTAKT_4}}", ""),
        ("{{HEADER_KONTAKT_5}}", ""),
        ("{{HEADER_KONTAKT_6}}", "{{HEADER_KONTAKT_6}}"),
    ]
    for cell, (text, url) in zip(cells, values):
        if url:
            add_hyperlink(cell, text, url, size=7.8, color=TEXT)
        else:
            add_text(cell, text, size=7.8, color=TEXT, line=1.05)


def add_strengths(parent) -> None:
    table = parent.add_table(rows=1, cols=3)
    set_table_geometry(table, [3_408, 3_408, 3_408])
    remove_table_borders(table)
    for index, cell in enumerate([cell for row in table.rows for cell in row.cells], 1):
        add_text(
            cell,
            f"{{{{STAERKE_{index}_TITEL}}}}",
            size=9.7,
            color=ACCENT,
            bold=True,
            after=1.5,
            keep_next=True,
        )
        add_text(
            cell,
            f"{{{{STAERKE_{index}_BESCHREIBUNG}}}}",
            size=8.2,
            color=TEXT,
            line=1.15,
        )


def add_career_entry(parent, index: int) -> None:
    table = parent.add_table(rows=1, cols=2)
    set_table_geometry(table, [8_200, 2_024])
    remove_table_borders(table)
    left, right = table.rows[0].cells
    add_text(
        left,
        f"{{{{POSITION_{index}}}}}",
        size=11.8,
        color=PRIMARY,
        after=0.5,
        keep_next=True,
    )
    add_text(
        left,
        f"{{{{UNTERNEHMEN_{index}}}}}",
        size=9.8,
        color=ACCENT,
        bold=True,
        after=1.2,
        keep_next=True,
    )
    add_text(
        left,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        size=8.3,
        color=TEXT,
        after=1.0,
        line=1.12,
    )
    for achievement in range(1, 6):
        add_real_bullet(left, f"{{{{ERFOLG_{index}_{achievement}}}}}")
    add_text(
        right,
        f"{{{{ARBEITSORT_{index}}}}}",
        size=7.8,
        color=MUTED,
        after=4,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    add_text(
        right,
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}{{{{ENDDATUM_{index}}}}}",
        size=7.8,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )


def add_education_entry(parent, index: int) -> None:
    table = parent.add_table(rows=1, cols=2)
    set_table_geometry(table, [8_200, 2_024])
    remove_table_borders(table)
    left, right = table.rows[0].cells
    add_text(
        left,
        f"{{{{ABSCHLUSS_{index}}}}}",
        size=11.3,
        color=PRIMARY,
        after=0.5,
        keep_next=True,
    )
    add_text(
        left,
        f"{{{{HOCHSCHULE_{index}}}}}",
        size=9.2,
        color=TEXT,
        after=1.5,
    )
    add_text(
        right,
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        size=7.8,
        color=MUTED,
        after=4,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    add_text(
        right,
        f"{{{{AUSBILDUNG_START_{index}}}}}{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}{{{{AUSBILDUNG_ENDE_{index}}}}}",
        size=7.8,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )


def build_visual(output: Path, tmp: Path) -> None:
    doc = Document()
    set_page(doc)
    configure_styles(doc, PRIMARY)
    background = tmp / "klassisch-background.png"
    photo = tmp / "klassisch-photo.png"
    draw_wave_background(background)
    create_placeholder_photo(photo)
    add_header_background(doc, background)

    header = doc.add_table(rows=1, cols=2)
    set_table_geometry(header, [8_290, 1_934])
    remove_table_borders(header)
    identity, photo_cell = header.rows[0].cells
    add_text(
        identity,
        "{{VORNAME}} {{NACHNAME}}",
        size=25.5,
        color=PRIMARY,
        bold=True,
        after=2,
        line=1.0,
        keep_next=True,
    )
    add_text(
        identity,
        "{{BERUFSBEZEICHNUNG}}{{#FACHGEBIETE}} | {{FACHGEBIETE}}{{/FACHGEBIETE}}",
        size=11.8,
        color=TEXT,
        after=2.5,
        line=1.08,
        keep_next=True,
    )
    add_contact_grid(identity)
    photo_paragraph = photo_cell.paragraphs[0]
    photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    photo_shape = photo_paragraph.add_run().add_picture(
        str(photo), width=Mm(34), height=Mm(34)
    )
    set_photo_shape(photo_shape)

    add_section_title(doc, "{{ZUSAMMENFASSUNG_TITEL}}")
    add_text(doc, "{{ZUSAMMENFASSUNG}}", size=8.5, color=TEXT, line=1.18)
    add_section_title(doc, "{{STAERKEN_TITEL}}")
    add_strengths(doc)
    add_section_title(doc, "{{ERFAHRUNG_TITEL}}")
    for index in range(1, 9):
        add_career_entry(doc, index)
    add_section_title(doc, "{{AUSBILDUNG_TITEL}}")
    for index in range(1, 4):
        add_education_entry(doc, index)
    add_section_title(doc, "{{KENNTNISSE_TITEL}}")
    add_text(doc, "{{KENNTNISSE}}", size=8.5, color=TEXT)
    add_section_title(doc, "{{SPRACHEN_TITEL}}")
    language_table = doc.add_table(rows=1, cols=3)
    set_table_geometry(language_table, [3_408, 3_408, 3_408])
    remove_table_borders(language_table)
    for index, cell in enumerate([cell for row in language_table.rows for cell in row.cells], 1):
        paragraph = add_text(
            cell,
            f"{{{{#SPRACHE_{index}}}}}{{{{SPRACHE_{index}}}}}",
            size=9.0,
            color=PRIMARY,
            after=0,
        )
        run = paragraph.add_run(
            f"{{{{#SPRACHNIVEAU_{index}}}}}  ({{{{SPRACHNIVEAU_{index}}}}})"
            f"{{{{/SPRACHNIVEAU_{index}}}}}{{{{/SPRACHE_{index}}}}}"
        )
        set_run(run, 8.6, TEXT)
    add_section_title(doc, "{{ZERTIFIKATE_TITEL}}")
    add_text(doc, "{{ZERTIFIKATE}}", size=8.5, color=TEXT)
    add_footer(doc)
    doc.save(output)


def build_ats(output: Path) -> None:
    doc = Document()
    set_page(doc)
    configure_styles(doc, ATS_BLUE, body_size=9.0)
    add_text(doc, "{{VORNAME}} {{NACHNAME}}", size=20, color=ATS_BLUE, bold=True, after=2)
    add_text(doc, "{{BERUFSBEZEICHNUNG}}", size=10.5, color=ATS_BLUE, bold=True, after=2)
    add_text(doc, "{{KONTAKT_ZEILE_1}}", size=8.4, color=TEXT, after=0.5)
    add_text(doc, "{{KONTAKT_ZEILE_2}}", size=8.4, color=TEXT, after=0.5)
    add_text(doc, "{{KONTAKT_ZEILE_3}}", size=8.4, color=TEXT, after=4)
    add_section_title(doc, "{{ZUSAMMENFASSUNG_TITEL}}", ats=True)
    add_text(doc, "{{ZUSAMMENFASSUNG}}", size=9.0, color=TEXT, line=1.18)
    add_section_title(doc, "{{BERUFSERFAHRUNG_TITEL}}", ats=True)
    for index in range(1, 9):
        add_text(doc, f"{{{{POSITION_{index}}}}}", size=10.5, color=ATS_BLUE, bold=True, after=0.5, keep_next=True)
        add_text(doc, f"{{{{UNTERNEHMEN_{index}}}}}{{{{METADATA_TRENNER_{index}}}}}{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}{{{{ENDDATUM_{index}}}}}{{{{METADATA_TRENNER_{index}}}}}{{{{ARBEITSORT_{index}}}}}", size=8.4, color=TEXT, after=1.0, keep_next=True)
        add_text(doc, f"{{{{BESCHREIBUNG_{index}}}}}", size=8.7, color=TEXT, after=0.5)
        for achievement in range(1, 6):
            add_real_bullet(doc, f"{{{{ERFOLG_{index}_{achievement}}}}}", size=8.5)
    add_section_title(doc, "{{AUSBILDUNG_TITEL}}", ats=True)
    for index in range(1, 4):
        add_text(doc, f"{{{{ABSCHLUSS_{index}}}}}", size=10.0, color=ATS_BLUE, bold=True, after=0.5, keep_next=True)
        add_text(doc, f"{{{{HOCHSCHULE_{index}}}}}{{{{AUSBILDUNG_METADATA_TRENNER_{index}}}}}{{{{AUSBILDUNG_START_{index}}}}}{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}{{{{AUSBILDUNG_ENDE_{index}}}}}{{{{AUSBILDUNG_METADATA_TRENNER_{index}}}}}{{{{AUSBILDUNG_ORT_{index}}}}}", size=8.5, color=TEXT, after=2)
    add_section_title(doc, "{{KENNTNISSE_TITEL}}", ats=True)
    add_text(doc, "{{KENNTNISSE}}", size=9.0, color=TEXT)
    add_section_title(doc, "{{SPRACHEN_TITEL}}", ats=True)
    add_text(doc, "{{SPRACHEN_ATS}}", size=9.0, color=TEXT)
    add_section_title(doc, "{{STAERKEN_TITEL}}", ats=True)
    add_text(doc, "{{STAERKEN_ATS}}", size=9.0, color=TEXT)
    add_section_title(doc, "{{ZERTIFIKATE_TITEL}}", ats=True)
    add_text(doc, "{{ZERTIFIKATE}}", size=9.0, color=TEXT)
    add_footer(doc, ats=True)
    doc.save(output)


def load_font(size: int, bold=False):
    filename = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / filename
    return ImageFont.truetype(str(path), size)


def wrap_text(draw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def preview(output: Path) -> None:
    image = Image.new("RGB", (1240, 1754), "white")
    tmp_bg = output.with_suffix(".background.png")
    draw_wave_background(tmp_bg)
    image.alpha_composite(Image.open(tmp_bg).convert("RGBA")) if image.mode == "RGBA" else image.paste(Image.open(tmp_bg), (0, 0), Image.open(tmp_bg))
    tmp_bg.unlink()
    draw = ImageDraw.Draw(image)
    x = 88
    draw.text((x, 78), "JULIAN FISCHER", font=load_font(43, True), fill="#111111")
    draw.text((x, 132), "Oberkellner | Japanische Kultur | Teamführung", font=load_font(25), fill=f"#{TEXT}")
    draw.text((x, 170), "+49 30 12345678   julianfischer@web.de   linkedin.com/in/   München, Bayern", font=load_font(16), fill=f"#{TEXT}")
    draw.ellipse((1030, 78, 1178, 226), fill="#D8E4E6")
    draw.ellipse((1073, 103, 1135, 165), fill="#68747A")
    draw.ellipse((1052, 157, 1156, 245), fill="#68747A")
    y = 266

    def section(title: str):
        nonlocal y
        draw.text((x, y), title, font=load_font(19, True), fill=f"#{HEADING}")
        y += 37

    section("ZUSAMMENFASSUNG")
    summary = "Erfahrener Oberkellner mit über 7 Jahren Erfahrung in führenden Restaurants. Experte in der japanischen Kultur und Küche, mit bewährten Fähigkeiten in Teamführung und Kundenbeziehungsmanagement."
    for line in wrap_text(draw, summary, load_font(16), 1050):
        draw.text((x, y), line, font=load_font(16), fill=f"#{TEXT}")
        y += 23
    y += 34
    section("STÄRKEN")
    strengths = [
        ("Teamleitung", "Erfolgreiche Führung von Teams mit bis zu 30 Mitarbeitern."),
        ("Kommunikation", "Effektive Abstimmung mit Küchen- und Managementteams."),
        ("Konfliktmanagement", "Professionelle Lösung von Kundenbeschwerden."),
    ]
    for index, (title, detail) in enumerate(strengths):
        sx = x + index * 355
        draw.text((sx, y), title, font=load_font(18, True), fill=f"#{ACCENT}")
        sy = y + 28
        for line in wrap_text(draw, detail, load_font(15), 300):
            draw.text((sx, sy), line, font=load_font(15), fill=f"#{TEXT}")
            sy += 21
    y += 125
    section("ERFAHRUNG")
    jobs = [
        ("Stellvertretender Restaurantleiter", "Sushi Palace", "Düsseldorf", "2019 - 2023", ["Schulung von 25 neuen Mitarbeitern im Kundenservice.", "Optimierung der operativen Abläufe und des Inventarsystems.", "Führung eines Teams von 30 Mitarbeitern."]),
        ("Oberkellner", "Miso Master", "Berlin", "2016 - 2019", ["Einführung eines neuen Schulungsprogramms für Kellner.", "Koordination von Veranstaltungen mit messbarem Umsatzanstieg."]),
        ("Service Manager", "Tokyo Tisch", "Hamburg", "2013 - 2016", ["Umsetzung neuer Qualitätsstandards.", "Koordination mit Küchenleitung und Serviceteam."]),
    ]
    for title, company, city, dates, bullets in jobs:
        draw.text((x, y), title, font=load_font(22), fill=f"#{PRIMARY}")
        draw.text((x, y + 31), company, font=load_font(18, True), fill=f"#{ACCENT}")
        draw.text((1070, y), city, font=load_font(15), fill=f"#{MUTED}", anchor="ra")
        draw.text((1070, y + 31), dates, font=load_font(15), fill=f"#{MUTED}", anchor="ra")
        y += 60
        for bullet in bullets:
            draw.text((x + 6, y), "•", font=load_font(14), fill=f"#{MUTED}")
            draw.text((x + 26, y), bullet, font=load_font(15), fill=f"#{TEXT}")
            y += 21
        y += 22
    section("AUSBILDUNG")
    for degree, university, city, dates in [
        ("Master in Hotel- und Restaurantmanagement", "Technische Universität München", "München", "2011 - 2013"),
        ("Bachelor in Gastgewerbemanagement", "Universität Heidelberg", "Heidelberg", "2008 - 2011"),
    ]:
        draw.text((x, y), degree, font=load_font(21), fill=f"#{PRIMARY}")
        draw.text((x, y + 29), university, font=load_font(17), fill=f"#{TEXT}")
        draw.text((1070, y), city, font=load_font(15), fill=f"#{MUTED}", anchor="ra")
        draw.text((1070, y + 29), dates, font=load_font(15), fill=f"#{MUTED}", anchor="ra")
        y += 75
    section("SPRACHEN")
    draw.text((x, y), "Deutsch   (Muttersprache)", font=load_font(17), fill=f"#{PRIMARY}")
    draw.text((455, y), "Englisch   (Versiert)", font=load_font(17), fill=f"#{PRIMARY}")
    draw.text((x, 1690), "www.beispiel.de", font=load_font(13), fill=f"#{MUTED}")
    image.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path("public/templates"))
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    tmp = output_dir / ".klassisch-assets"
    tmp.mkdir(exist_ok=True)
    outputs = {
        "visual": output_dir / "Klassisch_Lebenslauf_Muster.docx",
        "ats": output_dir / "Klassisch_Lebenslauf_ATS.docx",
        "preview": output_dir / "Klassisch_Lebenslauf_Muster.preview.png",
    }
    if not args.force:
        existing = [path for path in outputs.values() if path.exists()]
        if existing:
            raise SystemExit(f"Refusing to overwrite: {existing[0]}")
    build_visual(outputs["visual"], tmp)
    build_ats(outputs["ats"])
    preview(outputs["preview"])
    for child in tmp.iterdir():
        child.unlink()
    tmp.rmdir()
    for path in outputs.values():
        print(path)


if __name__ == "__main__":
    main()
