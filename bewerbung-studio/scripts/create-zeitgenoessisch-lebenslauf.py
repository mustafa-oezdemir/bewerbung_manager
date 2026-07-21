#!/usr/bin/env python3
"""Create the bundled Zeitgenoessisch Word Lebenslauf templates.

The generated files are editable OOXML documents. The visual version uses a
31/69 A4 layout table; the ATS version uses a single text column.
Existing output files are never overwritten unless --force is supplied.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips


# compact_reference_guide with named "Zeitgenoessisch A4 CV" overrides.
# Header treatment: customer_pack adapted to the CV name/title block.
PAGE_WIDTH_DXA = 11_906
PAGE_HEIGHT_DXA = 16_838
PAGE_MARGIN_DXA = 454  # 8 mm
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - (2 * PAGE_MARGIN_DXA)
LEFT_COLUMN_DXA = 3_409
RIGHT_COLUMN_DXA = CONTENT_WIDTH_DXA - LEFT_COLUMN_DXA
PRIMARY = "0F5B4A"
ACCENT = "36B779"
SOFT_ACCENT = "CDEEDF"
TITLE_BACKGROUND = "9DDBB9"
INK = "374247"
MUTED = "69757A"
LIGHT = "EDF8F2"
WHITE = "FFFFFF"
FONT = "Arial"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(
    cell,
    *,
    top: int = 180,
    bottom: int = 180,
    start: int = 0,
    end: int = 0,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{name}"))
        if border is None:
            border = OxmlElement(f"w:{name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def apply_table_geometry(
    table,
    widths_dxa: Iterable[int],
    *,
    indent_dxa: int = 0,
) -> None:
    widths = list(widths_dxa)
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def set_run_font(
    run,
    *,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_paragraph_indents(paragraph, *, sidebar: bool = False) -> None:
    paragraph.paragraph_format.left_indent = Twips(300 if sidebar else 260)
    paragraph.paragraph_format.right_indent = Twips(300 if sidebar else 260)


def configure_style(
    doc: Document,
    name: str,
    *,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.08,
    base: str | None = None,
) -> None:
    styles = doc.styles
    style = (
        styles[name]
        if name in styles
        else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    if base:
        style.base_style = styles[base]
    style.font.name = FONT
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    styles = [
        ("CV Name", 26, PRIMARY, False, False, 0, 4, 1.0, None),
        ("CV Title", 10.5, PRIMARY, True, False, 0, 1, 1.0, None),
        ("CV Fields", 8.5, MUTED, False, False, 0, 4, 1.05, None),
        ("Section Left", 11, PRIMARY, True, False, 9, 4, 1.0, None),
        ("Entry Role", 9.5, INK, True, False, 1, 1, 1.0, None),
        ("Entry Company", 9, PRIMARY, True, False, 6, 1, 1.0, None),
        ("Entry Meta", 8, MUTED, False, False, 0, 2, 1.0, None),
        ("Body Compact", 8.5, INK, False, False, 0, 2, 1.08, None),
        ("Sidebar Heading", 9.5, PRIMARY, True, False, 10, 4, 1.0, None),
        ("Sidebar Body", 8.2, INK, False, False, 0, 3, 1.10, None),
        ("Sidebar Strong", 8.5, PRIMARY, True, False, 3, 0, 1.05, None),
        ("Sidebar Muted", 7.7, MUTED, False, False, 0, 4, 1.08, None),
        ("Optional", 8.5, INK, False, False, 0, 2, 1.08, None),
        ("Optional Sidebar", 8.2, INK, False, False, 0, 3, 1.10, None),
        ("Optional Heading", 11, PRIMARY, True, False, 9, 4, 1.0, None),
        ("Optional Sidebar Heading", 9.5, PRIMARY, True, False, 10, 4, 1.0, None),
        ("Optional Bullet", 8.3, INK, False, False, 0, 1, 1.05, "List Bullet"),
        ("Strength Bullet", 8.5, PRIMARY, True, False, 3, 1, 1.05, "List Bullet"),
        ("Contact Line", 7.8, INK, False, False, 0, 3, 1.08, None),
        ("Language Level", 7.7, MUTED, False, False, 0, 2, 1.05, None),
        ("Cell Terminator", 1, WHITE, False, False, 0, 0, 1.0, None),
        ("ATS Name", 24, PRIMARY, True, False, 0, 2, 1.0, None),
        ("ATS Section", 12, PRIMARY, True, False, 10, 4, 1.0, None),
    ]
    for (
        name,
        size,
        color,
        bold,
        italic,
        before,
        after,
        line,
        base,
    ) in styles:
        configure_style(
            doc,
            name,
            size=size,
            color=color,
            bold=bold,
            italic=italic,
            before=before,
            after=after,
            line=line,
            base=base,
        )

    for style_name in ("List Bullet", "Optional Bullet"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(8.3)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Twips(520)
        style.paragraph_format.first_line_indent = Twips(-240)
        style.paragraph_format.space_after = Pt(1)
        style.paragraph_format.line_spacing = 1.05
    strength_style = doc.styles["Strength Bullet"]
    strength_style.paragraph_format.left_indent = Twips(420)
    strength_style.paragraph_format.first_line_indent = Twips(-220)
    strength_style.paragraph_format.space_after = Pt(1)
    strength_style.paragraph_format.line_spacing = 1.05


def configure_page(doc: Document, *, ats: bool = False) -> None:
    section = doc.sections[0]
    section.page_width = Twips(PAGE_WIDTH_DXA)
    section.page_height = Twips(PAGE_HEIGHT_DXA)
    margin = Mm(15) if ats else Twips(PAGE_MARGIN_DXA)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Mm(4)
    section.footer_distance = Mm(5)

    section_pr = section._sectPr
    doc_grid = section_pr.find(qn("w:docGrid"))
    if doc_grid is None:
        doc_grid = OxmlElement("w:docGrid")
        section_pr.append(doc_grid)
    doc_grid.set(qn("w:linePitch"), "312")


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)
    cell.add_paragraph()


def add_paragraph(
    container,
    text: str,
    style: str,
    *,
    sidebar: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
    keep_with_next: bool = False,
    keep_together: bool = False,
) -> object:
    paragraph = container.add_paragraph(style=style)
    set_paragraph_indents(paragraph, sidebar=sidebar)
    paragraph.alignment = alignment
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = keep_together
    paragraph.add_run(text)
    return paragraph


def add_left_section_heading(cell, placeholder: str) -> None:
    paragraph = add_paragraph(
        cell,
        placeholder,
        "Optional Heading",
        keep_with_next=True,
    )
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), ACCENT)
    border.append(bottom)
    p_pr.append(border)


def add_sidebar_heading(cell, placeholder: str) -> None:
    paragraph = add_paragraph(
        cell,
        placeholder,
        "Optional Sidebar Heading",
        sidebar=True,
        keep_with_next=True,
    )
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), ACCENT)
    border.append(bottom)
    p_pr.append(border)


def add_experience(cell, index: int) -> None:
    company = add_paragraph(
        cell,
        "",
        "Entry Company",
        keep_with_next=True,
        keep_together=True,
    )
    set_paragraph_indents(company)
    company.paragraph_format.tab_stops.add_tab_stop(
        Twips(RIGHT_COLUMN_DXA - 840),
        WD_TAB_ALIGNMENT.RIGHT,
    )
    company.add_run(f"{{{{UNTERNEHMEN_{index}}}}}")
    company.add_run("\t")
    location_run = company.add_run(f"{{{{ARBEITSORT_{index}}}}}")
    set_run_font(location_run, size=8, color=MUTED)

    role = add_paragraph(
        cell,
        "",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    set_paragraph_indents(role)
    role.paragraph_format.tab_stops.add_tab_stop(
        Twips(RIGHT_COLUMN_DXA - 840),
        WD_TAB_ALIGNMENT.RIGHT,
    )
    role.add_run(f"{{{{POSITION_{index}}}}}")
    role.add_run("\t")
    date_run = role.add_run(
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}}"
    )
    set_run_font(date_run, size=8, color=MUTED)

    add_paragraph(
        cell,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        "Optional",
        keep_together=True,
    ).style = "Body Compact"
    for achievement_index in range(1, 5):
        add_paragraph(
            cell,
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            "Optional Bullet",
            keep_together=True,
        )


def add_education(cell, index: int) -> None:
    company = add_paragraph(
        cell,
        "",
        "Entry Company",
        keep_with_next=True,
        keep_together=True,
    )
    set_paragraph_indents(company)
    company.paragraph_format.tab_stops.add_tab_stop(
        Twips(RIGHT_COLUMN_DXA - 840),
        WD_TAB_ALIGNMENT.RIGHT,
    )
    company.add_run(f"{{{{HOCHSCHULE_{index}}}}}")
    company.add_run("\t")
    location_run = company.add_run(f"{{{{AUSBILDUNG_ORT_{index}}}}}")
    set_run_font(location_run, size=8, color=MUTED)

    degree = add_paragraph(
        cell,
        "",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    set_paragraph_indents(degree)
    degree.paragraph_format.tab_stops.add_tab_stop(
        Twips(RIGHT_COLUMN_DXA - 840),
        WD_TAB_ALIGNMENT.RIGHT,
    )
    degree.add_run(
        f"{{{{ABSCHLUSS_{index}}}}} {{{{FACHRICHTUNG_{index}}}}}",
    )
    degree.add_run("\t")
    date_run = degree.add_run(
        f"{{{{AUSBILDUNG_START_{index}}}}}"
        f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
        f"{{{{AUSBILDUNG_ENDE_{index}}}}}"
    )
    set_run_font(date_run, size=8, color=MUTED)


def create_profile_placeholder(path: Path) -> None:
    image = Image.new("RGB", (760, 760), "#E7F5ED")
    draw = ImageDraw.Draw(image)
    draw.ellipse((245, 120, 515, 390), fill="#7FAD98")
    draw.ellipse((135, 365, 625, 855), fill="#7FAD98")
    image.save(path, "PNG")


def create_profile_decoration(path: Path) -> None:
    image = Image.new("RGBA", (900, 900), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse((20, 90, 600, 760), fill="#CDEEDF")
    draw.ellipse((410, 10, 890, 540), fill="#9DDBB9")
    draw.rounded_rectangle(
        (370, 480, 850, 860),
        radius=150,
        fill="#36B779",
    )
    image.save(path, "PNG")


def convert_inline_to_background_anchor(shape) -> None:
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    position_h.append(align_h)
    anchor.append(position_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    offset_v = OxmlElement("wp:posOffset")
    offset_v.text = "-95000"
    position_v.append(offset_v)
    anchor.append(position_v)

    extent = inline.find(qn("wp:extent"))
    effect_extent = inline.find(qn("wp:effectExtent"))
    doc_pr = inline.find(qn("wp:docPr"))
    frame_pr = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    for element in (extent, effect_extent):
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element in (doc_pr, frame_pr, graphic):
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def add_profile_picture(
    cell,
    placeholder_path: Path,
    decoration_path: Path,
) -> None:
    paragraph = cell.add_paragraph()
    set_paragraph_indents(paragraph, sidebar=True)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(9)
    decoration_run = paragraph.add_run()
    decoration_shape = decoration_run.add_picture(
        str(decoration_path),
        width=Mm(55),
        height=Mm(55),
    )
    decoration_shape._inline.docPr.set(
        "descr",
        "Organische gruene Profildekoration",
    )
    convert_inline_to_background_anchor(decoration_shape)
    run = paragraph.add_run()
    shape = run.add_picture(
        str(placeholder_path),
        width=Mm(44),
        height=Mm(44),
    )
    shape._inline.docPr.set("descr", "PROFILFOTO")
    shape._inline.docPr.set("title", "PROFILFOTO")
    geometry = shape._inline.find(".//" + qn("a:prstGeom"))
    if geometry is not None:
        geometry.set("prst", "ellipse")


def remove_initial_cell_paragraph(cell) -> None:
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        paragraph = cell.paragraphs[0]
        paragraph._element.getparent().remove(paragraph._element)


def add_title_box(cell) -> None:
    paragraph = add_paragraph(
        cell,
        "{{BERUFSBEZEICHNUNG}}\n{{FACHGEBIETE}}",
        "CV Title",
        keep_with_next=True,
        keep_together=True,
    )
    paragraph.paragraph_format.left_indent = Twips(360)
    paragraph.paragraph_format.right_indent = Twips(360)
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), TITLE_BACKGROUND)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "2")
        border.set(qn("w:space"), "5")
        border.set(qn("w:color"), TITLE_BACKGROUND)
        borders.append(border)
    p_pr.append(borders)


def build_zeitgenoessisch(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    placeholder_path = output_path.parent / ".zeit-profile-placeholder.png"
    decoration_path = output_path.parent / ".zeit-profile-decoration.png"
    create_profile_placeholder(placeholder_path)
    create_profile_decoration(decoration_path)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)

    table = doc.add_table(rows=1, cols=2)
    set_table_borders_none(table)
    apply_table_geometry(table, [LEFT_COLUMN_DXA, RIGHT_COLUMN_DXA])
    row = table.rows[0]
    row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    left, right = row.cells
    set_cell_shading(left, WHITE)
    set_cell_shading(right, WHITE)
    set_cell_margins(left, top=120, bottom=120, start=40, end=160)
    set_cell_margins(right, top=120, bottom=120, start=220, end=40)
    remove_initial_cell_paragraph(left)
    remove_initial_cell_paragraph(right)

    add_profile_picture(left, placeholder_path, decoration_path)

    add_sidebar_heading(left, "{{KONTAKTE_TITEL}}")
    for placeholder in (
        "TELEFON_ZEILE",
        "EMAIL_ZEILE",
        "WEBSITE_ZEILE",
        "LINKEDIN_ZEILE",
        "ORT_ZEILE",
    ):
        add_paragraph(
            left,
            f"{{{{{placeholder}}}}}",
            "Contact Line",
            sidebar=True,
        )

    add_sidebar_heading(left, "{{STAERKEN_TITEL}}")
    for index in range(1, 4):
        add_paragraph(
            left,
            f"{{{{STAERKE_{index}_TITEL}}}}",
            "Strength Bullet",
            sidebar=True,
            keep_with_next=True,
        )
        add_paragraph(
            left,
            f"{{{{STAERKE_{index}_BESCHREIBUNG}}}}",
            "Optional Sidebar",
            sidebar=True,
        ).style = "Sidebar Muted"

    add_sidebar_heading(left, "{{SPRACHEN_TITEL}}")
    for index in range(1, 4):
        add_paragraph(
            left,
            f"{{{{SPRACHE_{index}}}}}",
            "Optional Sidebar",
            sidebar=True,
            keep_with_next=True,
        ).style = "Sidebar Strong"
        add_paragraph(
            left,
            f"{{{{SPRACHNIVEAU_{index}}}}}  "
            f"{{{{SPRACHE_{index}_PUNKTE}}}}",
            "Optional Sidebar",
            sidebar=True,
        ).style = "Language Level"

    add_sidebar_heading(left, "{{KENNTNISSE_TITEL}}")
    for index in range(1, 6):
        add_paragraph(
            left,
            f"{{{{KENNTNIS_KATEGORIE_{index}}}}}",
            "Optional Sidebar",
            sidebar=True,
            keep_with_next=True,
        ).style = "Sidebar Strong"
        add_paragraph(
            left,
            f"{{{{KENNTNIS_EINTRAEGE_{index}}}}}",
            "Optional Sidebar",
            sidebar=True,
        ).style = "Sidebar Muted"

    for title_key, content_key in (
        ("SOFTWARE_TITEL", "SOFTWARE"),
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
        ("ZUSATZANGABEN_TITEL", "ZUSATZANGABEN"),
        ("FUEHRERSCHEIN_TITEL", "FUEHRERSCHEIN"),
        ("INTERESSEN_TITEL", "INTERESSEN"),
    ):
        add_sidebar_heading(left, f"{{{{{title_key}}}}}")
        add_paragraph(
            left,
            f"{{{{{content_key}}}}}",
            "Optional Sidebar",
            sidebar=True,
        )

    add_paragraph(
        right,
        "{{VORNAME}} {{NACHNAME}}",
        "CV Name",
        keep_with_next=True,
    )
    add_title_box(right)

    add_left_section_heading(right, "{{ZUSAMMENFASSUNG_TITEL}}")
    add_paragraph(
        right,
        "{{ZUSAMMENFASSUNG}}",
        "Optional",
        keep_together=True,
    ).style = "Body Compact"

    add_left_section_heading(right, "{{BERUFSERFAHRUNG_TITEL}}")
    for index in range(1, 7):
        add_experience(right, index)

    add_left_section_heading(right, "{{AUSBILDUNG_TITEL}}")
    for index in range(1, 4):
        add_education(right, index)

    for title_key, content_key in (
        ("PROJEKTE_TITEL", "PROJEKTE"),
        ("WEITERBILDUNGEN_TITEL", "WEITERBILDUNGEN"),
        ("VEROEFFENTLICHUNGEN_TITEL", "VEROEFFENTLICHUNGEN"),
        ("EHRENAMT_TITEL", "EHRENAMT"),
    ):
        add_left_section_heading(right, f"{{{{{title_key}}}}}")
        add_paragraph(right, f"{{{{{content_key}}}}}", "Optional")

    left.add_paragraph(" ", style="Cell Terminator")
    right.add_paragraph(" ", style="Cell Terminator")

    doc.save(output_path)
    placeholder_path.unlink(missing_ok=True)
    decoration_path.unlink(missing_ok=True)
    return output_path


def add_ats_experience(doc: Document, index: int) -> None:
    add_paragraph(
        doc,
        f"{{{{POSITION_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{UNTERNEHMEN_{index}}}}} | "
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}} | {{{{ARBEITSORT_{index}}}}}",
        "Optional",
        keep_with_next=True,
        keep_together=True,
    ).style = "Entry Meta"
    add_paragraph(
        doc,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        "Optional",
        keep_together=True,
    )
    for achievement_index in range(1, 5):
        add_paragraph(
            doc,
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            "Optional Bullet",
            keep_together=True,
        )


def add_ats_education(doc: Document, index: int) -> None:
    add_paragraph(
        doc,
        f"{{{{ABSCHLUSS_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{FACHRICHTUNG_{index}}}}}",
        "Optional",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{HOCHSCHULE_{index}}}}} | "
        f"{{{{AUSBILDUNG_START_{index}}}}}"
        f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
        f"{{{{AUSBILDUNG_ENDE_{index}}}}} | "
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        "Optional",
        keep_together=True,
    ).style = "Entry Meta"


def build_ats(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc, ats=True)
    configure_styles(doc)

    add_paragraph(
        doc,
        "{{VORNAME}} {{NACHNAME}}",
        "ATS Name",
        keep_with_next=True,
    )
    add_paragraph(
        doc,
        "{{BERUFSBEZEICHNUNG}}",
        "CV Title",
        keep_with_next=True,
    )
    add_paragraph(doc, "{{FACHGEBIETE}}", "Optional")
    add_paragraph(doc, "{{KONTAKT_ZEILE_1}}", "Optional")
    add_paragraph(doc, "{{KONTAKT_ZEILE_2}}", "Optional")
    add_paragraph(doc, "{{KONTAKT_ZEILE_3}}", "Optional")

    add_paragraph(
        doc,
        "{{ZUSAMMENFASSUNG_TITEL}}",
        "Optional Heading",
        keep_with_next=True,
    ).style = "ATS Section"
    add_paragraph(doc, "{{ZUSAMMENFASSUNG}}", "Optional")

    add_paragraph(
        doc,
        "{{BERUFSERFAHRUNG_TITEL}}",
        "Optional Heading",
        keep_with_next=True,
    ).style = "ATS Section"
    for index in range(1, 7):
        add_ats_experience(doc, index)

    add_paragraph(
        doc,
        "{{AUSBILDUNG_TITEL}}",
        "Optional Heading",
        keep_with_next=True,
    ).style = "ATS Section"
    for index in range(1, 4):
        add_ats_education(doc, index)

    for title_key, content_key in (
        ("KENNTNISSE_TITEL", "KENNTNISSE"),
        ("SPRACHEN_TITEL", "SPRACHEN_ATS"),
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
        ("WEITERBILDUNGEN_TITEL", "WEITERBILDUNGEN"),
        ("PROJEKTE_TITEL", "PROJEKTE"),
        ("EHRENAMT_TITEL", "EHRENAMT"),
    ):
        add_paragraph(
            doc,
            f"{{{{{title_key}}}}}",
            "Optional Heading",
            keep_with_next=True,
        ).style = "ATS Section"
        add_paragraph(doc, f"{{{{{content_key}}}}}", "Optional")

    doc.save(output_path)
    return output_path


def build_preview(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (840, 1188), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    def font(size: int, bold: bool = False):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
        return ImageFont.load_default()

    left_x, divider_x, right_x = 44, 274, 306
    draw.line((divider_x, 40, divider_x, 1148), fill="#CDEEDF", width=2)
    draw.ellipse((58, 48, 244, 240), fill="#CDEEDF")
    draw.ellipse((120, 32, 264, 170), fill="#9DDBB9")
    draw.rounded_rectangle((130, 150, 258, 252), radius=40, fill="#36B779")
    draw.ellipse((82, 62, 230, 210), fill="#FFFFFF")
    draw.ellipse((90, 70, 222, 202), fill="#7FAD98")
    draw.ellipse((134, 94, 178, 138), fill="#E7F5ED")
    draw.ellipse((112, 136, 200, 220), fill="#E7F5ED")

    draw.text((right_x, 52), "LENA HOFFMANN", font=font(41), fill="#0F5B4A")
    draw.rounded_rectangle((right_x, 112, 784, 174), radius=14, fill="#9DDBB9")
    draw.text(
        (right_x + 18, 126),
        "FINANZBUCHHALTUNG | CONTROLLING",
        font=font(16, True),
        fill="#0F5B4A",
    )

    def heading(x: int, y: int, text: str, width: int) -> int:
        draw.text((x, y), text, font=font(16, True), fill="#0F5B4A")
        draw.line((x, y + 27, x + width, y + 27), fill="#36B779", width=3)
        return y + 42

    def text_lines(x: int, y: int, widths: list[int], color: str = "#69757A") -> int:
        for width in widths:
            draw.rounded_rectangle((x, y, x + width, y + 8), radius=4, fill=color)
            y += 17
        return y

    y_left = heading(left_x, 278, "KONTAKTE", 185)
    y_left = text_lines(left_x, y_left, [168, 184, 152, 176, 118])
    y_left = heading(left_x, y_left + 24, "STÄRKEN", 185)
    for label_width, detail_width in ((98, 174), (150, 166), (112, 180)):
        draw.ellipse((left_x, y_left + 3, left_x + 9, y_left + 12), fill="#36B779")
        y_left = text_lines(left_x + 17, y_left, [label_width], "#0F5B4A")
        y_left = text_lines(left_x + 17, y_left, [detail_width, detail_width - 22])
        y_left += 7
    y_left = heading(left_x, y_left + 16, "SPRACHEN", 185)
    draw.text((left_x, y_left), "DEUTSCH", font=font(13, True), fill="#0F5B4A")
    draw.text((left_x, y_left + 23), "Muttersprache   ●●●●●", font=font(12), fill="#69757A")
    draw.text((left_x, y_left + 62), "ENGLISCH", font=font(13, True), fill="#0F5B4A")
    draw.text((left_x, y_left + 85), "Versiert        ●●●●○", font=font(12), fill="#69757A")

    y_right = heading(right_x, 218, "ZUSAMMENFASSUNG", 476)
    y_right = text_lines(right_x, y_right, [466, 452, 470, 430, 458, 340])
    y_right = heading(right_x, y_right + 24, "ERFAHRUNG", 476)
    for company, role in (
        ("SIEMENS AG", "Senior Buchhalterin"),
        ("MUSTER FINANCE GMBH", "Finanzbuchhalterin"),
        ("NORDSTERN AG", "Sachbearbeiterin Finanzen"),
    ):
        draw.text((right_x, y_right), company, font=font(15, True), fill="#0F5B4A")
        draw.text((right_x, y_right + 23), role, font=font(13, True), fill="#374247")
        draw.text((676, y_right + 3), "Berlin", font=font(11), fill="#69757A")
        draw.text((676, y_right + 24), "2019 – 2023", font=font(11), fill="#69757A")
        y_right = text_lines(right_x, y_right + 51, [450, 424])
        for width in (428, 444, 392):
            draw.ellipse((right_x + 3, y_right + 2, right_x + 10, y_right + 9), fill="#36B779")
            y_right = text_lines(right_x + 18, y_right, [width])
        y_right += 15
    y_right = heading(right_x, y_right + 4, "AUSBILDUNG", 476)
    draw.text(
        (right_x, y_right),
        "LUDWIG-MAXIMILIANS-UNIVERSITÄT",
        font=font(14, True),
        fill="#0F5B4A",
    )
    draw.text(
        (right_x, y_right + 25),
        "Master in Finanzbuchhaltung",
        font=font(12, True),
        fill="#374247",
    )
    draw.text((680, y_right + 25), "2010 – 2012", font=font(11), fill="#69757A")

    image.save(output_path, "PNG", optimize=True)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create Zeitgenoessisch Lebenslauf Word templates",
    )
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--ats-output", type=Path)
    parser.add_argument("--preview-output", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    print(build_zeitgenoessisch(args.output.resolve(), force=args.force))
    if args.ats_output:
        print(build_ats(args.ats_output.resolve(), force=args.force))
    if args.preview_output:
        print(build_preview(args.preview_output.resolve(), force=args.force))


if __name__ == "__main__":
    main()
