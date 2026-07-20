#!/usr/bin/env python3
"""Create the bundled Kompakt Word Lebenslauf templates.

The visual document is an editable A4 DOCX with a fixed 60/40 content table,
small safe margins, real Word bullets and a low-contrast technical line
decoration. The ATS document uses a single text column.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips


# compact_reference_guide with named "Kompakt A4 high-density CV" overrides.
# Header treatment: customer_pack reduced to a compact left-aligned name block.
PAGE_WIDTH_DXA = 11_906
PAGE_HEIGHT_DXA = 16_838
VERTICAL_MARGIN_DXA = 624  # 11 mm
HORIZONTAL_MARGIN_DXA = 850  # 15 mm
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - (2 * HORIZONTAL_MARGIN_DXA)
LEFT_COLUMN_DXA = 6_124
RIGHT_COLUMN_DXA = CONTENT_WIDTH_DXA - LEFT_COLUMN_DXA
PRIMARY = "0A3485"
ACCENT = "FF6500"
INK = "374247"
MUTED = "69757A"
LINE = "C9D0D4"
DECORATIVE = "FFD8BF"
INACTIVE = "E2E5E7"
WHITE = "FFFFFF"
FONT = "Arial"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(
    cell,
    *,
    top: int = 80,
    bottom: int = 80,
    start: int = 100,
    end: int = 100,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{name}"))
        if border is None:
            border = OxmlElement(f"w:{name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def apply_table_geometry(
    table,
    widths_dxa: Iterable[int],
    *,
    indent_dxa: int = 0,
) -> None:
    widths = list(widths_dxa)
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def configure_style(
    doc: Document,
    name: str,
    *,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.02,
    base: str | None = None,
) -> None:
    style = (
        doc.styles[name]
        if name in doc.styles
        else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    if base:
        style.base_style = doc.styles[base]
    style.font.name = FONT
    r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:cs"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal_r_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    normal_r_fonts.set(qn("w:ascii"), FONT)
    normal_r_fonts.set(qn("w:hAnsi"), FONT)
    normal_r_fonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.02

    styles = [
        ("CV Name", 25, PRIMARY, False, False, 0, 1, 1.0, None),
        ("CV Title", 10.5, ACCENT, True, False, 0, 3, 1.0, None),
        ("Section Main", 10.5, PRIMARY, True, False, 5, 2, 1.0, None),
        ("Section Side", 10, PRIMARY, True, False, 6, 2, 1.0, None),
        ("Entry Role", 9.6, PRIMARY, True, False, 2, 0, 1.0, None),
        ("Entry Company", 9.2, ACCENT, True, False, 0, 0, 1.0, None),
        ("Entry Meta", 8.3, MUTED, False, False, 0, 1, 1.0, None),
        ("Entry Body", 9.0, INK, False, False, 0, 1, 1.02, None),
        ("Entry Tech", 8.4, PRIMARY, False, True, 0, 1, 1.0, None),
        ("Contact Link", 9.0, INK, False, False, 0, 1, 1.0, None),
        ("Contact Body", 9.0, INK, False, False, 0, 1, 1.0, None),
        ("Side Strong", 9.0, PRIMARY, True, False, 2, 0, 1.0, None),
        ("Side Body", 9.0, INK, False, False, 0, 1, 1.02, None),
        ("Side Muted", 9.0, MUTED, False, False, 0, 1, 1.0, None),
        ("Language Name", 9.0, PRIMARY, True, False, 0, 0, 1.0, None),
        ("Language Level", 8.0, MUTED, False, False, 0, 0, 1.0, None),
        ("Skill Tag", 9.0, PRIMARY, True, False, 1, 0, 1.0, None),
        ("Skill Value", 9.0, INK, False, False, 0, 1, 1.0, None),
        ("Highlight Title", 9.0, PRIMARY, True, False, 2, 0, 1.0, None),
        ("Highlight Body", 9.0, MUTED, False, False, 0, 1, 1.0, None),
        ("Optional", 9.0, INK, False, False, 0, 1, 1.02, None),
        ("Optional Bullet", 9.0, INK, False, False, 0, 0, 1.02, "List Bullet"),
        ("Cell Terminator", 1, WHITE, False, False, 0, 0, 1.0, None),
        ("ATS Name", 24, PRIMARY, True, False, 0, 2, 1.0, None),
        ("ATS Title", 10.5, ACCENT, True, False, 0, 2, 1.0, None),
        ("ATS Section", 12, PRIMARY, True, False, 9, 3, 1.0, None),
    ]
    for name, size, color, bold, italic, before, after, line, base in styles:
        configure_style(
            doc,
            name,
            size=size,
            color=color,
            bold=bold,
            italic=italic,
            before=before,
            after=after,
            line=line,
            base=base,
        )
    for style_name in ("List Bullet", "Optional Bullet"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(9.0)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Twips(430)
        style.paragraph_format.first_line_indent = Twips(-210)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.02


def configure_page(doc: Document, *, ats: bool = False) -> None:
    section = doc.sections[0]
    section.page_width = Twips(PAGE_WIDTH_DXA)
    section.page_height = Twips(PAGE_HEIGHT_DXA)
    if ats:
        section.top_margin = Mm(15)
        section.bottom_margin = Mm(15)
        section.left_margin = Mm(17)
        section.right_margin = Mm(17)
    else:
        section.top_margin = Twips(VERTICAL_MARGIN_DXA)
        section.bottom_margin = Twips(VERTICAL_MARGIN_DXA)
        section.left_margin = Twips(HORIZONTAL_MARGIN_DXA)
        section.right_margin = Twips(HORIZONTAL_MARGIN_DXA)
    section.header_distance = Mm(3)
    section.footer_distance = Mm(4)
    doc_grid = section._sectPr.find(qn("w:docGrid"))
    if doc_grid is None:
        doc_grid = OxmlElement("w:docGrid")
        section._sectPr.append(doc_grid)
    doc_grid.set(qn("w:linePitch"), "300")


def add_paragraph(
    container,
    text: str,
    style: str,
    *,
    keep_with_next: bool = False,
    keep_together: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
) -> object:
    paragraph = container.add_paragraph(style=style)
    paragraph.alignment = alignment
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = keep_together
    paragraph.add_run(text)
    return paragraph


def remove_initial_cell_paragraph(cell) -> None:
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        paragraph = cell.paragraphs[0]
        paragraph._element.getparent().remove(paragraph._element)


def add_heading(container, placeholder: str, *, side: bool = False) -> None:
    paragraph = add_paragraph(
        container,
        placeholder,
        "Section Side" if side else "Section Main",
        keep_with_next=True,
    )
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5" if side else "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    p_pr.append(borders)


def add_experience(container, index: int) -> None:
    add_paragraph(
        container,
        f"{{{{POSITION_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        container,
        f"{{{{UNTERNEHMEN_{index}}}}}",
        "Entry Company",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        container,
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}} {{{{METADATA_TRENNER_{index}}}}} "
        f"{{{{ARBEITSORT_{index}}}}}",
        "Entry Meta",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        container,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        "Entry Body",
        keep_together=True,
    )
    for achievement_index in range(1, 6):
        add_paragraph(
            container,
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            "Optional Bullet",
            keep_together=True,
        )
    add_paragraph(
        container,
        f"{{{{TECHNOLOGIEN_{index}}}}}",
        "Entry Tech",
        keep_together=True,
    )


def add_education(container, index: int) -> None:
    add_paragraph(
        container,
        f"{{{{ABSCHLUSS_{index}}}}} {{{{FACHRICHTUNG_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        container,
        f"{{{{HOCHSCHULE_{index}}}}}",
        "Entry Company",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        container,
        f"{{{{AUSBILDUNG_START_{index}}}}}"
        f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
        f"{{{{AUSBILDUNG_ENDE_{index}}}}} "
        f"{{{{AUSBILDUNG_METADATA_TRENNER_{index}}}}} "
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        "Entry Meta",
        keep_together=True,
    )


def create_technical_pattern(path: Path) -> None:
    image = Image.new("RGBA", (1100, 900), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    color = (255, 216, 191, 128)
    paths = [
        [(80, 160), (300, 160), (410, 260), (700, 260), (850, 120), (1030, 120)],
        [(20, 350), (240, 350), (390, 500), (650, 500), (810, 340), (1060, 340)],
        [(140, 720), (360, 720), (500, 590), (730, 590), (900, 760), (1080, 760)],
    ]
    for points in paths:
        draw.line(points, fill=color, width=5, joint="curve")
        for x, y in (points[0], points[-1], points[2], points[-2]):
            draw.ellipse((x - 11, y - 11, x + 11, y + 11), outline=color, width=5)
    image.save(path, "PNG", optimize=True)


def convert_inline_to_background_anchor(shape) -> None:
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    position_h.append(align_h)
    anchor.append(position_h)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    offset_v = OxmlElement("wp:posOffset")
    offset_v.text = "-120000"
    position_v.append(offset_v)
    anchor.append(position_v)
    for element_name in ("wp:extent", "wp:effectExtent"):
        element = inline.find(qn(element_name))
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element_name in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        element = inline.find(qn(element_name))
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def add_background_pattern(cell, pattern_path: Path) -> None:
    paragraph = cell.add_paragraph(style="Cell Terminator")
    run = paragraph.add_run()
    shape = run.add_picture(str(pattern_path), width=Mm(67), height=Mm(55))
    shape._inline.docPr.set("descr", "KOMPAKT_DEKORATION")
    shape._inline.docPr.set("title", "KOMPAKT_DEKORATION")
    convert_inline_to_background_anchor(shape)


def add_languages(cell) -> None:
    add_heading(cell, "{{SPRACHEN_TITEL}}")
    table = cell.add_table(rows=1, cols=3)
    set_table_borders_none(table)
    apply_table_geometry(table, [1_900, 1_900, 1_900])
    for index, language_cell in enumerate(table.rows[0].cells, start=1):
        set_cell_margins(language_cell, top=20, bottom=20, start=0, end=80)
        remove_initial_cell_paragraph(language_cell)
        add_paragraph(
            language_cell,
            f"{{{{SPRACHE_{index}}}}}",
            "Language Name",
            keep_with_next=True,
        )
        add_paragraph(
            language_cell,
            f"{{{{SPRACHNIVEAU_{index}}}}}",
            "Language Level",
            keep_with_next=True,
        )
        paragraph = add_paragraph(
            language_cell,
            f"{{{{SPRACHE_{index}_PUNKTE}}}}",
            "Language Level",
        )
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(ACCENT)
        language_cell.add_paragraph(" ", style="Cell Terminator")


def add_skills(cell) -> None:
    add_heading(cell, "{{KENNTNISSE_TITEL}}", side=True)
    table = cell.add_table(rows=3, cols=2)
    set_table_borders_none(table)
    apply_table_geometry(table, [1_850, 1_850])
    for index, skill_cell in enumerate(
        [item for row in table.rows for item in row.cells],
        start=1,
    ):
        set_cell_margins(skill_cell, top=20, bottom=25, start=0, end=90)
        remove_initial_cell_paragraph(skill_cell)
        add_paragraph(
            skill_cell,
            f"{{{{KENNTNIS_KATEGORIE_{index}}}}}",
            "Skill Tag",
            keep_with_next=True,
        )
        add_paragraph(
            skill_cell,
            f"{{{{KENNTNIS_EINTRAEGE_{index}}}}}",
            "Skill Value",
        )
        skill_cell.add_paragraph(" ", style="Cell Terminator")


def build_kompakt(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pattern_path = output_path.parent / ".kompakt-technical-pattern.png"
    create_technical_pattern(pattern_path)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)

    add_paragraph(
        doc,
        "{{VORNAME}} {{NACHNAME}}",
        "CV Name",
        keep_with_next=True,
    )
    add_paragraph(
        doc,
        "{{BERUFSBEZEICHNUNG}}",
        "CV Title",
        keep_with_next=True,
    )

    content = doc.add_table(rows=1, cols=2)
    set_table_borders_none(content)
    apply_table_geometry(content, [LEFT_COLUMN_DXA, RIGHT_COLUMN_DXA])
    left, right = content.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(left, WHITE)
    set_cell_shading(right, WHITE)
    set_cell_margins(left, top=20, bottom=20, start=0, end=160)
    set_cell_margins(right, top=20, bottom=20, start=180, end=0)
    remove_initial_cell_paragraph(left)
    remove_initial_cell_paragraph(right)

    add_heading(left, "{{ERFAHRUNG_TITEL}}")
    for index in range(1, 9):
        add_experience(left, index)
    add_heading(left, "{{AUSBILDUNG_TITEL}}")
    for index in range(1, 4):
        add_education(left, index)
    add_languages(left)

    add_background_pattern(right, pattern_path)
    add_heading(right, "{{KONTAKTDATEN_TITEL}}", side=True)
    for key, style in (
        ("TELEFON", "Contact Link"),
        ("EMAIL", "Contact Link"),
        ("LINKEDIN", "Contact Link"),
        ("GITHUB", "Contact Link"),
        ("WEBSITE", "Contact Link"),
        ("ORT", "Contact Body"),
        ("GEBURTSZEILE", "Contact Body"),
    ):
        add_paragraph(right, f"{{{{{key}}}}}", style, keep_together=True)

    add_heading(right, "{{ZUSAMMENFASSUNG_TITEL}}", side=True)
    add_paragraph(
        right,
        "{{ZUSAMMENFASSUNG}}",
        "Side Body",
        keep_together=True,
    )
    add_heading(right, "{{STAERKEN_TITEL}}", side=True)
    for index in range(1, 4):
        add_paragraph(
            right,
            f"{{{{STAERKE_{index}_TITEL}}}}",
            "Side Strong",
            keep_with_next=True,
        )
        add_paragraph(
            right,
            f"{{{{STAERKE_{index}_BESCHREIBUNG}}}}",
            "Side Muted",
        )

    add_heading(right, "{{ERFOLGE_TITEL}}", side=True)
    for index in range(1, 3):
        add_paragraph(
            right,
            f"{{{{ERFOLG_HIGHLIGHT_{index}_TITEL}}}}",
            "Highlight Title",
            keep_with_next=True,
        )
        add_paragraph(
            right,
            f"{{{{ERFOLG_HIGHLIGHT_{index}_BESCHREIBUNG}}}}",
            "Highlight Body",
        )
    add_skills(right)
    left.add_paragraph(" ", style="Cell Terminator")
    right.add_paragraph(" ", style="Cell Terminator")
    doc.save(output_path)
    pattern_path.unlink(missing_ok=True)
    return output_path


def add_ats_experience(doc: Document, index: int) -> None:
    add_paragraph(
        doc,
        f"{{{{POSITION_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{UNTERNEHMEN_{index}}}}} | "
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}} | {{{{ARBEITSORT_{index}}}}}",
        "Entry Meta",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        "Entry Body",
        keep_together=True,
    )
    for achievement_index in range(1, 6):
        add_paragraph(
            doc,
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            "Optional Bullet",
            keep_together=True,
        )
    add_paragraph(
        doc,
        f"{{{{TECHNOLOGIEN_{index}}}}}",
        "Entry Tech",
        keep_together=True,
    )


def add_ats_education(doc: Document, index: int) -> None:
    add_paragraph(
        doc,
        f"{{{{ABSCHLUSS_{index}}}}} {{{{FACHRICHTUNG_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{HOCHSCHULE_{index}}}}} | "
        f"{{{{AUSBILDUNG_START_{index}}}}}"
        f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
        f"{{{{AUSBILDUNG_ENDE_{index}}}}} | "
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        "Entry Meta",
        keep_together=True,
    )


def build_ats(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc, ats=True)
    configure_styles(doc)
    add_paragraph(doc, "{{VORNAME}} {{NACHNAME}}", "ATS Name", keep_with_next=True)
    add_paragraph(doc, "{{BERUFSBEZEICHNUNG}}", "ATS Title", keep_with_next=True)
    for key in ("TELEFON", "EMAIL", "LINKEDIN", "GITHUB", "WEBSITE", "ORT", "GEBURTSZEILE"):
        add_paragraph(doc, f"{{{{{key}}}}}", "Contact Body")
    add_paragraph(
        doc,
        "{{ZUSAMMENFASSUNG_TITEL}}",
        "ATS Section",
        keep_with_next=True,
    )
    add_paragraph(doc, "{{ZUSAMMENFASSUNG}}", "Optional")
    add_paragraph(doc, "{{ERFAHRUNG_TITEL}}", "ATS Section", keep_with_next=True)
    for index in range(1, 9):
        add_ats_experience(doc, index)
    add_paragraph(doc, "{{AUSBILDUNG_TITEL}}", "ATS Section", keep_with_next=True)
    for index in range(1, 4):
        add_ats_education(doc, index)
    for title_key, content_key in (
        ("SPRACHEN_TITEL", "SPRACHEN_ATS"),
        ("STAERKEN_TITEL", "STAERKEN_ATS"),
        ("ERFOLGE_TITEL", "ERFOLGE_ATS"),
        ("KENNTNISSE_TITEL", "KENNTNISSE"),
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
    ):
        add_paragraph(
            doc,
            f"{{{{{title_key}}}}}",
            "ATS Section",
            keep_with_next=True,
        )
        add_paragraph(doc, f"{{{{{content_key}}}}}", "Optional")
    doc.save(output_path)
    return output_path


def build_preview(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (840, 1188), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    def font(size: int, bold: bool = False):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
        return ImageFont.load_default()

    def heading(x: int, y: int, text: str, width: int) -> int:
        draw.text((x, y), text, font=font(14, True), fill="#0A3485")
        draw.line((x, y + 24, x + width, y + 24), fill="#FF6500", width=2)
        return y + 34

    def bars(x: int, y: int, widths: list[int], color: str = "#69757A") -> int:
        for width in widths:
            draw.rounded_rectangle((x, y, x + width, y + 6), radius=3, fill=color)
            y += 12
        return y

    draw.text((54, 38), "JULIAN FISCHER", font=font(36), fill="#0A3485")
    draw.text((56, 86), "IT-PROJEKTMANAGER", font=font(14, True), fill="#FF6500")
    left_x, divider_x, right_x = 54, 500, 526
    draw.line((divider_x, 124, divider_x, 1140), fill="#C9D0D4", width=1)
    pattern = [
        [(650, 138), (720, 138), (750, 168), (816, 168)],
        [(690, 190), (748, 190), (778, 220), (830, 220)],
        [(624, 268), (704, 268), (738, 234), (814, 234)],
    ]
    for points in pattern:
        draw.line(points, fill="#FFD8BF", width=2)
        for x, y in (points[0], points[-1]):
            draw.ellipse((x - 4, y - 4, x + 4, y + 4), outline="#FFD8BF", width=2)

    y_left = heading(left_x, 132, "ERFAHRUNG", 420)
    for role, company in (
        ("Stellv. Restaurantleiter", "Sushi Palace"),
        ("Schichtleiter", "Urban Food GmbH"),
        ("Servicekoordinator", "City Catering AG"),
        ("Teamleiter", "Gastro Service KG"),
    ):
        draw.text((left_x, y_left), role, font=font(12, True), fill="#0A3485")
        draw.text((left_x, y_left + 19), company, font=font(11, True), fill="#FF6500")
        draw.text((left_x, y_left + 38), "2019-2023 · Düsseldorf", font=font(9), fill="#69757A")
        y_left = bars(left_x, y_left + 56, [402, 374])
        for width in (366, 390, 344):
            draw.ellipse((left_x + 2, y_left, left_x + 7, y_left + 5), fill="#FF6500")
            y_left = bars(left_x + 13, y_left, [width])
        y_left += 10
    y_left = heading(left_x, y_left + 2, "AUSBILDUNG", 420)
    draw.text((left_x, y_left), "B.A. Betriebswirtschaft", font=font(11, True), fill="#0A3485")
    draw.text((left_x, y_left + 19), "Hochschule Düsseldorf", font=font(10, True), fill="#FF6500")
    draw.text((left_x, y_left + 38), "2014-2018 · Düsseldorf", font=font(9), fill="#69757A")
    y_left = heading(left_x, y_left + 66, "SPRACHEN", 420)
    for offset, name, level in (
        (0, "DEUTSCH", "Muttersprache  ●●●●●"),
        (140, "ENGLISCH", "Versiert  ●●●●○"),
        (280, "JAPANISCH", "Grundkenntnisse  ●●○○○"),
    ):
        draw.text((left_x + offset, y_left), name, font=font(9, True), fill="#0A3485")
        draw.text((left_x + offset, y_left + 18), level, font=font(8), fill="#69757A")

    y_right = heading(right_x, 132, "KONTAKTDATEN", 260)
    y_right = bars(right_x, y_right, [222, 250, 244, 188, 232, 160])
    y_right = heading(right_x, y_right + 14, "ZUSAMMENFASSUNG", 260)
    y_right = bars(right_x, y_right, [250, 238, 254, 224, 246, 194])
    y_right = heading(right_x, y_right + 14, "STÄRKEN", 260)
    for width in (122, 152, 134):
        y_right = bars(right_x, y_right, [width], "#0A3485")
        y_right = bars(right_x, y_right, [244, 210])
        y_right += 5
    y_right = heading(right_x, y_right + 10, "ERFOLGE", 260)
    for width in (144, 168):
        y_right = bars(right_x, y_right, [width], "#0A3485")
        y_right = bars(right_x, y_right, [242, 204])
        y_right += 5
    y_right = heading(right_x, y_right + 10, "FÄHIGKEITEN", 260)
    for row in range(3):
        draw.rounded_rectangle(
            (right_x, y_right + row * 45, right_x + 116, y_right + 28 + row * 45),
            radius=7,
            outline="#FFD8BF",
            width=2,
        )
        draw.rounded_rectangle(
            (right_x + 130, y_right + row * 45, right_x + 252, y_right + 28 + row * 45),
            radius=7,
            outline="#FFD8BF",
            width=2,
        )
    image.save(output_path, "PNG", optimize=True)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create Kompakt Lebenslauf Word templates",
    )
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--ats-output", type=Path)
    parser.add_argument("--preview-output", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    print(build_kompakt(args.output.resolve(), force=args.force))
    if args.ats_output:
        print(build_ats(args.ats_output.resolve(), force=args.force))
    if args.preview_output:
        print(build_preview(args.preview_output.resolve(), force=args.force))


if __name__ == "__main__":
    main()
