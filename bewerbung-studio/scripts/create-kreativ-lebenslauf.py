#!/usr/bin/env python3
"""Create the bundled Kreativ Word Lebenslauf templates.

The generated files are editable OOXML documents. The visual version uses a
full-width header banner plus a compact 60/40 A4 content table. The ATS
version uses a single text column.
Existing output files are never overwritten unless --force is supplied.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips


# compact_reference_guide with named "Kreativ A4 one-page CV" overrides.
# Header treatment: customer_pack adapted to a full-width banner.
PAGE_WIDTH_DXA = 11_906
PAGE_HEIGHT_DXA = 16_838
PAGE_MARGIN_DXA = 454  # 8 mm
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - (2 * PAGE_MARGIN_DXA)
HEADER_LEFT_DXA = 8_250
HEADER_RIGHT_DXA = CONTENT_WIDTH_DXA - HEADER_LEFT_DXA
LEFT_COLUMN_DXA = 6_599
RIGHT_COLUMN_DXA = CONTENT_WIDTH_DXA - LEFT_COLUMN_DXA
PRIMARY = "154F45"
ACCENT = "39B774"
HEADER_BACKGROUND = "39B774"
HEADER_TEXT = "FFFFFF"
INK = "39454A"
MUTED = "69757A"
LIGHT = "E4EBE8"
LINE = "C8D0D2"
WHITE = "FFFFFF"
FONT = "Arial"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(
    cell,
    *,
    top: int = 180,
    bottom: int = 180,
    start: int = 0,
    end: int = 0,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{name}"))
        if border is None:
            border = OxmlElement(f"w:{name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def apply_table_geometry(
    table,
    widths_dxa: Iterable[int],
    *,
    indent_dxa: int = 0,
) -> None:
    widths = list(widths_dxa)
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def set_run_font(
    run,
    *,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_paragraph_indents(paragraph, *, sidebar: bool = False) -> None:
    paragraph.paragraph_format.left_indent = Twips(180 if sidebar else 200)
    paragraph.paragraph_format.right_indent = Twips(180 if sidebar else 200)


def configure_style(
    doc: Document,
    name: str,
    *,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.08,
    base: str | None = None,
) -> None:
    styles = doc.styles
    style = (
        styles[name]
        if name in styles
        else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    if base:
        style.base_style = styles[base]
    style.font.name = FONT
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(8.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.02

    styles = [
        ("CV Name", 24, HEADER_TEXT, True, False, 0, 2, 1.0, None),
        ("CV Title", 9.5, HEADER_TEXT, True, False, 0, 2, 1.0, None),
        ("CV Fields", 7.6, HEADER_TEXT, False, False, 0, 1, 1.0, None),
        ("Section Left", 10.5, PRIMARY, True, False, 7, 3, 1.0, None),
        ("Entry Role", 9.2, INK, True, False, 2, 0, 1.0, None),
        ("Entry Company", 8.5, ACCENT, True, False, 0, 0, 1.0, None),
        ("Entry Meta", 7.6, MUTED, False, False, 0, 1, 1.0, None),
        ("Entry Technologies", 7.5, PRIMARY, False, True, 0, 1, 1.0, None),
        ("Entry Separator", 1, WHITE, False, False, 1, 2, 1.0, None),
        ("Body Compact", 8.2, INK, False, False, 0, 1, 1.02, None),
        ("Sidebar Heading", 9.5, PRIMARY, True, False, 10, 4, 1.0, None),
        ("Sidebar Body", 8.1, INK, False, False, 0, 2, 1.04, None),
        ("Sidebar Strong", 8.2, PRIMARY, True, False, 2, 0, 1.0, None),
        ("Sidebar Muted", 7.7, MUTED, False, False, 0, 2, 1.02, None),
        ("Optional", 8.2, INK, False, False, 0, 1, 1.02, None),
        ("Optional Sidebar", 8.1, INK, False, False, 0, 2, 1.04, None),
        ("Optional Heading", 10.5, PRIMARY, True, False, 7, 3, 1.0, None),
        ("Optional Sidebar Heading", 9.5, PRIMARY, True, False, 7, 3, 1.0, None),
        ("Optional Bullet", 8.1, INK, False, False, 0, 0, 1.02, "List Bullet"),
        ("Strength Bullet", 8.2, PRIMARY, True, False, 2, 0, 1.02, "List Bullet"),
        ("Contact Line", 7.6, HEADER_TEXT, False, False, 0, 1, 1.0, None),
        ("Language Level", 7.7, MUTED, False, False, 0, 1, 1.0, None),
        ("Cell Terminator", 1, WHITE, False, False, 0, 0, 1.0, None),
        ("ATS Name", 24, PRIMARY, True, False, 0, 2, 1.0, None),
        ("ATS Title", 10.5, ACCENT, True, False, 0, 2, 1.0, None),
        ("ATS Section", 12, PRIMARY, True, False, 10, 4, 1.0, None),
    ]
    for (
        name,
        size,
        color,
        bold,
        italic,
        before,
        after,
        line,
        base,
    ) in styles:
        configure_style(
            doc,
            name,
            size=size,
            color=color,
            bold=bold,
            italic=italic,
            before=before,
            after=after,
            line=line,
            base=base,
        )

    for style_name in ("List Bullet", "Optional Bullet"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(8.1)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Twips(520)
        style.paragraph_format.first_line_indent = Twips(-240)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.02
    strength_style = doc.styles["Strength Bullet"]
    strength_style.paragraph_format.left_indent = Twips(420)
    strength_style.paragraph_format.first_line_indent = Twips(-220)
    strength_style.paragraph_format.space_after = Pt(0)
    strength_style.paragraph_format.line_spacing = 1.02


def configure_page(doc: Document, *, ats: bool = False) -> None:
    section = doc.sections[0]
    section.page_width = Twips(PAGE_WIDTH_DXA)
    section.page_height = Twips(PAGE_HEIGHT_DXA)
    margin = Mm(15) if ats else Twips(PAGE_MARGIN_DXA)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Mm(4)
    section.footer_distance = Mm(5)

    section_pr = section._sectPr
    doc_grid = section_pr.find(qn("w:docGrid"))
    if doc_grid is None:
        doc_grid = OxmlElement("w:docGrid")
        section_pr.append(doc_grid)
    doc_grid.set(qn("w:linePitch"), "312")


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)
    cell.add_paragraph()


def add_paragraph(
    container,
    text: str,
    style: str,
    *,
    sidebar: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
    keep_with_next: bool = False,
    keep_together: bool = False,
) -> object:
    paragraph = container.add_paragraph(style=style)
    set_paragraph_indents(paragraph, sidebar=sidebar)
    paragraph.alignment = alignment
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = keep_together
    paragraph.add_run(text)
    return paragraph


def add_left_section_heading(cell, placeholder: str) -> None:
    paragraph = add_paragraph(
        cell,
        placeholder,
        "Optional Heading",
        keep_with_next=True,
    )
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), ACCENT)
    border.append(bottom)
    p_pr.append(border)


def add_sidebar_heading(cell, placeholder: str) -> None:
    paragraph = add_paragraph(
        cell,
        placeholder,
        "Optional Sidebar Heading",
        sidebar=True,
        keep_with_next=True,
    )
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), ACCENT)
    border.append(bottom)
    p_pr.append(border)


def add_entry_separator(cell, index: int) -> None:
    paragraph = add_paragraph(
        cell,
        f"{{{{ERFAHRUNG_TRENNER_{index}}}}}",
        "Entry Separator",
        keep_together=True,
    )
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "dashed")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LINE)
    borders.append(bottom)
    p_pr.append(borders)


def add_experience(cell, index: int) -> None:
    add_paragraph(
        cell,
        f"{{{{POSITION_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        cell,
        f"{{{{UNTERNEHMEN_{index}}}}}",
        "Entry Company",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        cell,
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}} {{{{METADATA_TRENNER_{index}}}}} "
        f"{{{{ARBEITSORT_{index}}}}}",
        "Entry Meta",
        keep_with_next=True,
        keep_together=True,
    )

    add_paragraph(
        cell,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        "Optional",
        keep_together=True,
    ).style = "Body Compact"
    for achievement_index in range(1, 5):
        add_paragraph(
            cell,
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            "Optional Bullet",
            keep_together=True,
        )
    add_paragraph(
        cell,
        f"{{{{TECHNOLOGIEN_{index}}}}}",
        "Entry Technologies",
        keep_together=True,
    )
    add_entry_separator(cell, index)


def add_education(cell, index: int) -> None:
    add_paragraph(
        cell,
        f"{{{{ABSCHLUSS_{index}}}}} {{{{FACHRICHTUNG_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        cell,
        f"{{{{HOCHSCHULE_{index}}}}}",
        "Entry Company",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        cell,
        f"{{{{AUSBILDUNG_START_{index}}}}}"
        f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
        f"{{{{AUSBILDUNG_ENDE_{index}}}}} "
        f"{{{{AUSBILDUNG_METADATA_TRENNER_{index}}}}} "
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        "Entry Meta",
        keep_together=True,
    )


def create_profile_placeholder(path: Path) -> None:
    image = Image.new("RGB", (760, 760), "#DCE8E4")
    draw = ImageDraw.Draw(image)
    draw.ellipse((245, 120, 515, 390), fill="#82A89F")
    draw.ellipse((135, 365, 625, 855), fill="#82A89F")
    image.save(path, "PNG")


def create_background_pattern(path: Path) -> None:
    image = Image.new("RGBA", (1200, 1200), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    for offset, width in ((40, 10), (150, 8), (270, 7), (400, 6)):
        draw.ellipse(
            (offset, offset - 80, 1220 - offset, 1100 - offset),
            outline=(228, 235, 232, 145),
            width=width,
        )
    image.save(path, "PNG")


def convert_inline_to_background_anchor(shape) -> None:
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    position_h.append(align_h)
    anchor.append(position_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    offset_v = OxmlElement("wp:posOffset")
    offset_v.text = "-95000"
    position_v.append(offset_v)
    anchor.append(position_v)

    extent = inline.find(qn("wp:extent"))
    effect_extent = inline.find(qn("wp:effectExtent"))
    doc_pr = inline.find(qn("wp:docPr"))
    frame_pr = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    for element in (extent, effect_extent):
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element in (doc_pr, frame_pr, graphic):
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def add_profile_picture(cell, placeholder_path: Path) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run()
    shape = run.add_picture(
        str(placeholder_path),
        width=Mm(34),
        height=Mm(34),
    )
    shape._inline.docPr.set("descr", "PROFILFOTO")
    shape._inline.docPr.set("title", "PROFILFOTO")
    geometry = shape._inline.find(".//" + qn("a:prstGeom"))
    if geometry is not None:
        geometry.set("prst", "roundRect")


def add_background_pattern(cell, pattern_path: Path) -> None:
    paragraph = cell.add_paragraph(style="Cell Terminator")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run()
    shape = run.add_picture(
        str(pattern_path),
        width=Mm(68),
        height=Mm(68),
    )
    shape._inline.docPr.set(
        "descr",
        "Sehr dezentes dekoratives Kreislinienmuster",
    )
    convert_inline_to_background_anchor(shape)


def remove_initial_cell_paragraph(cell) -> None:
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        paragraph = cell.paragraphs[0]
        paragraph._element.getparent().remove(paragraph._element)


def add_header_contacts(cell) -> None:
    table = cell.add_table(rows=3, cols=2)
    set_table_borders_none(table)
    apply_table_geometry(table, [3_900, 3_900])
    placeholders = [
        "HEADER_KONTAKT_1",
        "HEADER_KONTAKT_2",
        "HEADER_KONTAKT_3",
        "HEADER_KONTAKT_4",
        "HEADER_KONTAKT_5",
        "HEADER_KONTAKT_6",
    ]
    for nested_cell, placeholder in zip(
        [cell for row in table.rows for cell in row.cells],
        placeholders,
    ):
        set_cell_shading(nested_cell, HEADER_BACKGROUND)
        set_cell_margins(
            nested_cell,
            top=15,
            bottom=15,
            start=20,
            end=80,
        )
        remove_initial_cell_paragraph(nested_cell)
        add_paragraph(
            nested_cell,
            f"{{{{{placeholder}}}}}",
            "Contact Line",
        )
        nested_cell.add_paragraph(" ", style="Cell Terminator")


def build_kreativ(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    placeholder_path = output_path.parent / ".kreativ-profile-placeholder.png"
    pattern_path = output_path.parent / ".kreativ-background-pattern.png"
    create_profile_placeholder(placeholder_path)
    create_background_pattern(pattern_path)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)

    banner = doc.add_table(rows=1, cols=2)
    set_table_borders_none(banner)
    apply_table_geometry(
        banner,
        [HEADER_LEFT_DXA, HEADER_RIGHT_DXA],
    )
    banner_left, banner_right = banner.rows[0].cells
    banner_left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    banner_right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for banner_cell in (banner_left, banner_right):
        set_cell_shading(banner_cell, HEADER_BACKGROUND)
    set_cell_margins(
        banner_left,
        top=150,
        bottom=130,
        start=220,
        end=80,
    )
    set_cell_margins(
        banner_right,
        top=90,
        bottom=90,
        start=60,
        end=180,
    )
    remove_initial_cell_paragraph(banner_left)
    remove_initial_cell_paragraph(banner_right)
    add_paragraph(
        banner_left,
        "{{VORNAME}} {{NACHNAME}}",
        "CV Name",
        keep_with_next=True,
    )
    add_paragraph(
        banner_left,
        "{{BERUFSBEZEICHNUNG}} | {{FACHGEBIETE}}",
        "CV Title",
        keep_with_next=True,
    )
    add_header_contacts(banner_left)
    add_profile_picture(banner_right, placeholder_path)

    spacer = doc.add_paragraph(style="Cell Terminator")
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(1)

    content = doc.add_table(rows=1, cols=2)
    set_table_borders_none(content)
    apply_table_geometry(
        content,
        [LEFT_COLUMN_DXA, RIGHT_COLUMN_DXA],
    )
    left, right = content.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(left, WHITE)
    set_cell_shading(right, WHITE)
    set_cell_margins(left, top=60, bottom=60, start=40, end=150)
    set_cell_margins(right, top=60, bottom=60, start=190, end=40)
    remove_initial_cell_paragraph(left)
    remove_initial_cell_paragraph(right)

    add_left_section_heading(left, "{{ERFAHRUNG_TITEL}}")
    for index in range(1, 7):
        add_experience(left, index)

    add_left_section_heading(left, "{{AUSBILDUNG_TITEL}}")
    for index in range(1, 4):
        add_education(left, index)

    for title_key, content_key in (
        ("PROJEKTE_TITEL", "PROJEKTE"),
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
        ("WEITERBILDUNGEN_TITEL", "WEITERBILDUNGEN"),
        ("VEROEFFENTLICHUNGEN_TITEL", "VEROEFFENTLICHUNGEN"),
        ("EHRENAMT_TITEL", "EHRENAMT"),
    ):
        add_left_section_heading(left, f"{{{{{title_key}}}}}")
        add_paragraph(left, f"{{{{{content_key}}}}}", "Optional")

    add_background_pattern(right, pattern_path)
    add_sidebar_heading(right, "{{ZUSAMMENFASSUNG_TITEL}}")
    add_paragraph(
        right,
        "{{ZUSAMMENFASSUNG}}",
        "Optional Sidebar",
        sidebar=True,
        keep_together=True,
    )

    add_sidebar_heading(right, "{{STAERKEN_TITEL}}")
    for index in range(1, 4):
        add_paragraph(
            right,
            f"{{{{STAERKE_{index}_TITEL}}}}",
            "Strength Bullet",
            sidebar=True,
            keep_with_next=True,
        )
        add_paragraph(
            right,
            f"{{{{STAERKE_{index}_BESCHREIBUNG}}}}",
            "Optional Sidebar",
            sidebar=True,
        ).style = "Sidebar Muted"

    add_sidebar_heading(right, "{{SPRACHEN_TITEL}}")
    for index in range(1, 4):
        add_paragraph(
            right,
            f"{{{{SPRACHE_{index}}}}}",
            "Optional Sidebar",
            sidebar=True,
            keep_with_next=True,
        ).style = "Sidebar Strong"
        add_paragraph(
            right,
            f"{{{{SPRACHNIVEAU_{index}}}}}  "
            f"{{{{SPRACHE_{index}_PUNKTE}}}}",
            "Optional Sidebar",
            sidebar=True,
        ).style = "Language Level"

    add_sidebar_heading(right, "{{KENNTNISSE_TITEL}}")
    for index in range(1, 6):
        add_paragraph(
            right,
            f"{{{{KENNTNIS_KATEGORIE_{index}}}}}",
            "Optional Sidebar",
            sidebar=True,
            keep_with_next=True,
        ).style = "Sidebar Strong"
        add_paragraph(
            right,
            f"{{{{KENNTNIS_EINTRAEGE_{index}}}}}",
            "Optional Sidebar",
            sidebar=True,
        ).style = "Sidebar Muted"

    left.add_paragraph(" ", style="Cell Terminator")
    right.add_paragraph(" ", style="Cell Terminator")

    doc.save(output_path)
    placeholder_path.unlink(missing_ok=True)
    pattern_path.unlink(missing_ok=True)
    return output_path


def add_ats_experience(doc: Document, index: int) -> None:
    add_paragraph(
        doc,
        f"{{{{POSITION_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{UNTERNEHMEN_{index}}}}} | "
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}} | {{{{ARBEITSORT_{index}}}}}",
        "Optional",
        keep_with_next=True,
        keep_together=True,
    ).style = "Entry Meta"
    add_paragraph(
        doc,
        f"{{{{BESCHREIBUNG_{index}}}}}",
        "Optional",
        keep_together=True,
    )
    for achievement_index in range(1, 5):
        add_paragraph(
            doc,
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            "Optional Bullet",
            keep_together=True,
        )
    add_paragraph(
        doc,
        f"{{{{TECHNOLOGIEN_{index}}}}}",
        "Entry Technologies",
        keep_together=True,
    )


def add_ats_education(doc: Document, index: int) -> None:
    add_paragraph(
        doc,
        f"{{{{ABSCHLUSS_{index}}}}}",
        "Entry Role",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{FACHRICHTUNG_{index}}}}}",
        "Optional",
        keep_with_next=True,
        keep_together=True,
    )
    add_paragraph(
        doc,
        f"{{{{HOCHSCHULE_{index}}}}} | "
        f"{{{{AUSBILDUNG_START_{index}}}}}"
        f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
        f"{{{{AUSBILDUNG_ENDE_{index}}}}} | "
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        "Optional",
        keep_together=True,
    ).style = "Entry Meta"


def build_ats(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc, ats=True)
    configure_styles(doc)

    add_paragraph(
        doc,
        "{{VORNAME}} {{NACHNAME}}",
        "ATS Name",
        keep_with_next=True,
    )
    add_paragraph(
        doc,
        "{{BERUFSBEZEICHNUNG}}",
        "ATS Title",
        keep_with_next=True,
    )
    add_paragraph(doc, "{{FACHGEBIETE}}", "Optional")
    add_paragraph(doc, "{{KONTAKT_ZEILE_1}}", "Optional")
    add_paragraph(doc, "{{KONTAKT_ZEILE_2}}", "Optional")
    add_paragraph(doc, "{{KONTAKT_ZEILE_3}}", "Optional")

    add_paragraph(
        doc,
        "{{ZUSAMMENFASSUNG_TITEL}}",
        "Optional Heading",
        keep_with_next=True,
    ).style = "ATS Section"
    add_paragraph(doc, "{{ZUSAMMENFASSUNG}}", "Optional")

    add_paragraph(
        doc,
        "{{ERFAHRUNG_TITEL}}",
        "Optional Heading",
        keep_with_next=True,
    ).style = "ATS Section"
    for index in range(1, 7):
        add_ats_experience(doc, index)

    add_paragraph(
        doc,
        "{{AUSBILDUNG_TITEL}}",
        "Optional Heading",
        keep_with_next=True,
    ).style = "ATS Section"
    for index in range(1, 4):
        add_ats_education(doc, index)

    for title_key, content_key in (
        ("KENNTNISSE_TITEL", "KENNTNISSE"),
        ("SPRACHEN_TITEL", "SPRACHEN_ATS"),
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
        ("WEITERBILDUNGEN_TITEL", "WEITERBILDUNGEN"),
        ("PROJEKTE_TITEL", "PROJEKTE"),
        ("EHRENAMT_TITEL", "EHRENAMT"),
    ):
        add_paragraph(
            doc,
            f"{{{{{title_key}}}}}",
            "Optional Heading",
            keep_with_next=True,
        ).style = "ATS Section"
        add_paragraph(doc, f"{{{{{content_key}}}}}", "Optional")

    doc.save(output_path)
    return output_path


def build_preview(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (840, 1188), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    def font(size: int, bold: bool = False):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
        return ImageFont.load_default()

    left_x, divider_x, right_x = 44, 274, 306
    draw.line((divider_x, 40, divider_x, 1148), fill="#CDEEDF", width=2)
    draw.ellipse((58, 48, 244, 240), fill="#CDEEDF")
    draw.ellipse((120, 32, 264, 170), fill="#9DDBB9")
    draw.rounded_rectangle((130, 150, 258, 252), radius=40, fill="#36B779")
    draw.ellipse((82, 62, 230, 210), fill="#FFFFFF")
    draw.ellipse((90, 70, 222, 202), fill="#7FAD98")
    draw.ellipse((134, 94, 178, 138), fill="#E7F5ED")
    draw.ellipse((112, 136, 200, 220), fill="#E7F5ED")

    draw.text((right_x, 52), "LENA HOFFMANN", font=font(41), fill="#0F5B4A")
    draw.rounded_rectangle((right_x, 112, 784, 174), radius=14, fill="#9DDBB9")
    draw.text(
        (right_x + 18, 126),
        "FINANZBUCHHALTUNG | CONTROLLING",
        font=font(16, True),
        fill="#0F5B4A",
    )

    def heading(x: int, y: int, text: str, width: int) -> int:
        draw.text((x, y), text, font=font(16, True), fill="#0F5B4A")
        draw.line((x, y + 27, x + width, y + 27), fill="#36B779", width=3)
        return y + 42

    def text_lines(x: int, y: int, widths: list[int], color: str = "#69757A") -> int:
        for width in widths:
            draw.rounded_rectangle((x, y, x + width, y + 8), radius=4, fill=color)
            y += 17
        return y

    y_left = heading(left_x, 278, "KONTAKTE", 185)
    y_left = text_lines(left_x, y_left, [168, 184, 152, 176, 118])
    y_left = heading(left_x, y_left + 24, "STÄRKEN", 185)
    for label_width, detail_width in ((98, 174), (150, 166), (112, 180)):
        draw.ellipse((left_x, y_left + 3, left_x + 9, y_left + 12), fill="#36B779")
        y_left = text_lines(left_x + 17, y_left, [label_width], "#0F5B4A")
        y_left = text_lines(left_x + 17, y_left, [detail_width, detail_width - 22])
        y_left += 7
    y_left = heading(left_x, y_left + 16, "SPRACHEN", 185)
    draw.text((left_x, y_left), "DEUTSCH", font=font(13, True), fill="#0F5B4A")
    draw.text((left_x, y_left + 23), "Muttersprache   ●●●●●", font=font(12), fill="#69757A")
    draw.text((left_x, y_left + 62), "ENGLISCH", font=font(13, True), fill="#0F5B4A")
    draw.text((left_x, y_left + 85), "Versiert        ●●●●○", font=font(12), fill="#69757A")

    y_right = heading(right_x, 218, "ZUSAMMENFASSUNG", 476)
    y_right = text_lines(right_x, y_right, [466, 452, 470, 430, 458, 340])
    y_right = heading(right_x, y_right + 24, "ERFAHRUNG", 476)
    for company, role in (
        ("SIEMENS AG", "Senior Buchhalterin"),
        ("MUSTER FINANCE GMBH", "Finanzbuchhalterin"),
        ("NORDSTERN AG", "Sachbearbeiterin Finanzen"),
    ):
        draw.text((right_x, y_right), company, font=font(15, True), fill="#0F5B4A")
        draw.text((right_x, y_right + 23), role, font=font(13, True), fill="#374247")
        draw.text((676, y_right + 3), "Berlin", font=font(11), fill="#69757A")
        draw.text((676, y_right + 24), "2019 – 2023", font=font(11), fill="#69757A")
        y_right = text_lines(right_x, y_right + 51, [450, 424])
        for width in (428, 444, 392):
            draw.ellipse((right_x + 3, y_right + 2, right_x + 10, y_right + 9), fill="#36B779")
            y_right = text_lines(right_x + 18, y_right, [width])
        y_right += 15
    y_right = heading(right_x, y_right + 4, "AUSBILDUNG", 476)
    draw.text(
        (right_x, y_right),
        "LUDWIG-MAXIMILIANS-UNIVERSITÄT",
        font=font(14, True),
        fill="#0F5B4A",
    )
    draw.text(
        (right_x, y_right + 25),
        "Master in Finanzbuchhaltung",
        font=font(12, True),
        fill="#374247",
    )
    draw.text((680, y_right + 25), "2010 – 2012", font=font(11), fill="#69757A")

    image.save(output_path, "PNG", optimize=True)
    return output_path


def build_kreativ_preview(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (840, 1188), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    def font(size: int, bold: bool = False):
        candidates = [
            Path(
                "C:/Windows/Fonts/arialbd.ttf"
                if bold
                else "C:/Windows/Fonts/arial.ttf",
            ),
            Path(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
                if bold
                else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            ),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
        return ImageFont.load_default()

    def heading(x: int, y: int, text: str, width: int) -> int:
        draw.text((x, y), text, font=font(15, True), fill="#154F45")
        draw.line((x, y + 25, x + width, y + 25), fill="#39B774", width=2)
        return y + 37

    def text_lines(
        x: int,
        y: int,
        widths: list[int],
        color: str = "#69757A",
    ) -> int:
        for width in widths:
            draw.rounded_rectangle(
                (x, y, x + width, y + 7),
                radius=3,
                fill=color,
            )
            y += 14
        return y

    draw.rectangle((0, 0, 840, 220), fill="#39B774")
    draw.text((42, 32), "MARIE SCHRÖDER", font=font(38, True), fill="#FFFFFF")
    draw.text(
        (44, 84),
        "IT-PROZESS- UND PROJEKTMANAGERIN | SOFTWAREENTWICKLUNG",
        font=font(13, True),
        fill="#FFFFFF",
    )
    draw.text((44, 126), "+49 30 12345678", font=font(12), fill="#FFFFFF")
    draw.text(
        (280, 126),
        "marie.schroeder@example.de",
        font=font(12),
        fill="#FFFFFF",
    )
    draw.text(
        (44, 153),
        "linkedin.com/in/marie",
        font=font(12),
        fill="#FFFFFF",
    )
    draw.text((280, 153), "Berlin, Deutschland", font=font(12), fill="#FFFFFF")
    draw.text(
        (44, 180),
        "Geb. 01.03.1990 in München",
        font=font(12),
        fill="#FFFFFF",
    )
    draw.rounded_rectangle((678, 32, 796, 188), radius=12, fill="#FFFFFF")
    draw.rounded_rectangle((686, 40, 788, 180), radius=10, fill="#DCE8E4")
    draw.ellipse((718, 58, 756, 96), fill="#82A89F")
    draw.ellipse((704, 96, 770, 172), fill="#82A89F")

    left_x, divider_x, right_x = 42, 510, 536
    draw.line((divider_x, 244, divider_x, 1148), fill="#C8D0D2", width=2)
    for offset in (0, 36, 72):
        draw.ellipse(
            (650 + offset, 248 + offset, 910 + offset, 508 + offset),
            outline="#E4EBE8",
            width=3,
        )

    y_left = heading(left_x, 252, "ERFAHRUNG", 430)
    for role, company in (
        ("Senior IT-Projektmanagerin", "TechnologieLösungen AG"),
        ("IT-Projektmanagerin", "Digital Systems GmbH"),
        ("Softwareentwicklerin", "Innovatech AG"),
    ):
        draw.text((left_x, y_left), role, font=font(13, True), fill="#39454A")
        draw.text(
            (left_x, y_left + 21),
            company,
            font=font(12, True),
            fill="#39B774",
        )
        draw.text(
            (left_x, y_left + 42),
            "2018–2023 · München",
            font=font(10),
            fill="#69757A",
        )
        y_left = text_lines(left_x, y_left + 64, [410, 388])
        for width in (382, 402, 356):
            draw.ellipse(
                (left_x + 3, y_left + 1, left_x + 9, y_left + 7),
                fill="#39B774",
            )
            y_left = text_lines(left_x + 15, y_left, [width])
        draw.line(
            (left_x, y_left + 5, left_x + 430, y_left + 5),
            fill="#C8D0D2",
            width=1,
        )
        y_left += 18
    y_left = heading(left_x, y_left + 2, "AUSBILDUNG", 430)
    draw.text(
        (left_x, y_left),
        "M.Sc. Wirtschaftsinformatik",
        font=font(12, True),
        fill="#39454A",
    )
    draw.text(
        (left_x, y_left + 22),
        "Technische Universität München",
        font=font(11, True),
        fill="#39B774",
    )
    draw.text(
        (left_x, y_left + 43),
        "2010–2012 · München",
        font=font(10),
        fill="#69757A",
    )

    y_right = heading(right_x, 252, "ZUSAMMENFASSUNG", 262)
    y_right = text_lines(right_x, y_right, [252, 244, 256, 228, 248, 186])
    y_right = heading(right_x, y_right + 18, "STÄRKEN", 262)
    for label_width, detail_width in ((112, 242), (92, 232), (128, 246)):
        draw.ellipse(
            (right_x, y_right + 2, right_x + 7, y_right + 9),
            fill="#39B774",
        )
        y_right = text_lines(
            right_x + 14,
            y_right,
            [label_width],
            "#154F45",
        )
        y_right = text_lines(
            right_x + 14,
            y_right,
            [detail_width, detail_width - 28],
        )
        y_right += 5
    y_right = heading(right_x, y_right + 12, "SPRACHEN", 262)
    draw.text((right_x, y_right), "DEUTSCH", font=font(11, True), fill="#154F45")
    draw.text(
        (right_x, y_right + 20),
        "Muttersprache     ●●●●●",
        font=font(10),
        fill="#69757A",
    )
    draw.text(
        (right_x, y_right + 48),
        "ENGLISCH",
        font=font(11, True),
        fill="#154F45",
    )
    draw.text(
        (right_x, y_right + 68),
        "Fortgeschritten   ●●●○○",
        font=font(10),
        fill="#69757A",
    )
    y_right = heading(right_x, y_right + 105, "FÄHIGKEITEN", 262)
    for label, width in (
        ("Projektmanagement", 218),
        ("Jira · Asana", 174),
        ("KI · Full Stack", 202),
    ):
        draw.text(
            (right_x, y_right),
            label,
            font=font(10, True),
            fill="#154F45",
        )
        y_right = text_lines(right_x, y_right + 20, [width])
        y_right += 7

    image.save(output_path, "PNG", optimize=True)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create Kreativ Lebenslauf Word templates",
    )
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--ats-output", type=Path)
    parser.add_argument("--preview-output", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    print(build_kreativ(args.output.resolve(), force=args.force))
    if args.ats_output:
        print(build_ats(args.ats_output.resolve(), force=args.force))
    if args.preview_output:
        print(build_kreativ_preview(args.preview_output.resolve(), force=args.force))


if __name__ == "__main__":
    main()
