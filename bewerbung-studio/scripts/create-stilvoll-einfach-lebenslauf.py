#!/usr/bin/env python3
"""Create the bundled Stilvoll and Einfach Word Lebenslauf templates."""

from __future__ import annotations

import argparse
import math
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

A4_WIDTH_MM = 210
A4_HEIGHT_MM = 297
FONT = "Arial"
INK = "424B50"
MUTED = "667177"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size: float, color: str = INK, bold: bool = False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def configure_paragraph(
    paragraph,
    *,
    before: float = 0,
    after: float = 0,
    line: float = 1.0,
    keep_next: bool = False,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def add_text(
    parent,
    text: str,
    *,
    size: float = 8.5,
    color: str = INK,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.15,
    keep_next: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    paragraph = parent.add_paragraph()
    paragraph.alignment = align
    configure_paragraph(
        paragraph,
        before=before,
        after=after,
        line=line,
        keep_next=keep_next,
    )
    set_run(paragraph.add_run(text), size, color, bold)
    return paragraph


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_bottom_border(paragraph, color: str, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_section_title(
    parent,
    title: str,
    *,
    color: str,
    size: float,
    border_color: str | None = None,
    uppercase: bool = True,
) -> None:
    paragraph = add_text(
        parent,
        title.upper() if uppercase else title,
        size=size,
        color=color,
        bold=True,
        before=7,
        after=5,
        line=1.0,
        keep_next=True,
    )
    if border_color:
        add_bottom_border(paragraph, border_color, 10)


def add_real_bullet(parent, text: str, *, size=8.2, color=INK) -> None:
    paragraph = parent.add_paragraph(style="List Bullet")
    configure_paragraph(paragraph, after=1, line=1.1)
    paragraph.paragraph_format.left_indent = Mm(4)
    paragraph.paragraph_format.first_line_indent = Mm(-2)
    set_run(paragraph.add_run(text), size, color)


def set_page(doc: Document, margins: tuple[float, float, float, float]) -> None:
    section = doc.sections[0]
    section.page_width = Mm(A4_WIDTH_MM)
    section.page_height = Mm(A4_HEIGHT_MM)
    top, right, bottom, left = margins
    section.top_margin = Mm(top)
    section.right_margin = Mm(right)
    section.bottom_margin = Mm(bottom)
    section.left_margin = Mm(left)
    section.header_distance = Mm(3)
    section.footer_distance = Mm(5)


def configure_styles(doc: Document, body_size: float, primary: str) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    bullet = doc.styles["List Bullet"]
    bullet.font.name = FONT
    bullet.font.size = Pt(body_size)
    bullet.font.color.rgb = rgb(INK)
    bullet.paragraph_format.left_indent = Mm(4)
    bullet.paragraph_format.first_line_indent = Mm(-2)
    bullet.paragraph_format.space_after = Pt(1)
    bullet.paragraph_format.line_spacing = 1.1
    for style_name, size in (("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 9.5)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(primary)
        style.paragraph_format.keep_with_next = True


def set_photo_alt_text(inline_shape, alt_text="PROFILFOTO") -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)


def create_placeholder_photo(path: Path, palette: tuple[str, str]) -> None:
    image = Image.new("RGB", (420, 420), f"#{palette[1]}")
    draw = ImageDraw.Draw(image)
    draw.ellipse((80, 45, 340, 305), fill=f"#{palette[0]}")
    draw.rounded_rectangle((55, 250, 365, 470), 80, fill=f"#{palette[0]}")
    image.save(path)


def draw_stilvoll_pattern(path: Path) -> None:
    image = Image.new("RGBA", (1240, 1754), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    color = (220, 226, 223, 72)
    for x in range(420, 1220, 210):
        for y in range(-60, 720, 140):
            draw.line((x - 105, y + 105, x, y, x + 105, y + 105), fill=color, width=2)
            draw.line((x - 105, y + 130, x, y + 25, x + 105, y + 130), fill=color, width=2)
    image.save(path)


def draw_einfach_pattern(path: Path) -> None:
    image = Image.new("RGBA", (1240, 1754), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    color = (234, 246, 253, 220)
    paths = [
        [(755, -20), (755, 80), (845, 80), (845, -20)],
        [(260, 140), (260, 260), (350, 260), (350, 190)],
        [(910, 120), (910, 250), (1010, 250), (1010, 190)],
        [(680, 1530), (680, 1700), (770, 1700)],
        [(50, 1590), (50, 1680)],
    ]
    for points in paths:
        draw.line(points, fill=color, width=28, joint="curve")
    draw.arc((20, 1520, 160, 1660), 90, 300, fill=color, width=28)
    image.save(path)


def add_header_background(doc: Document, image_path: Path) -> None:
    header = doc.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Mm(210), height=Mm(297))
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", "0"),
        ("behindDoc", "1"),
        ("locked", "0"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
    ):
        anchor.set(key, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for tag in ("positionH", "positionV"):
        position = OxmlElement(f"wp:{tag}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    for element in (
        inline.find(qn("wp:extent")),
        inline.find(qn("wp:effectExtent")),
    ):
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element in (
        inline.find(qn("wp:docPr")),
        inline.find(qn("wp:cNvGraphicFramePr")),
        inline.find(qn("a:graphic")),
    ):
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def add_footer(doc: Document, color: str) -> None:
    section = doc.sections[0]
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Mm(180))
    set_table_geometry(table, [5100, 5100])
    remove_table_borders(table)
    left, right = table.rows[0].cells
    add_text(left, "{{WEBSITE}}", size=7.2, color=color)
    paragraph = add_text(
        right,
        "Seite ",
        size=7.2,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([fld_begin, instr, separate, text, end])


def add_career_entry(
    parent,
    index: int,
    *,
    primary: str,
    accent: str,
    compact: bool = False,
) -> None:
    add_text(
        parent,
        f"{{{{POSITION_{index}}}}}",
        size=10.5 if compact else 11,
        color=primary,
        bold=not compact,
        before=2,
        after=1,
        keep_next=True,
    )
    paragraph = parent.add_paragraph()
    configure_paragraph(paragraph, after=2, line=1.05, keep_next=True)
    set_run(paragraph.add_run(f"{{{{UNTERNEHMEN_{index}}}}}"), 9.3, accent, True)
    set_run(
        paragraph.add_run(
            f"    {{{{STARTDATUM_{index}}}}}–{{{{ENDDATUM_{index}}}}}    {{{{ARBEITSORT_{index}}}}}"
        ),
        7.8,
        MUTED,
    )
    for achievement in range(1, 5):
        add_real_bullet(
            parent,
            f"{{{{ERFOLG_{index}_{achievement}}}}}",
            size=7.9 if compact else 8.2,
        )


def build_stilvoll_visual(output: Path, tmp: Path) -> None:
    doc = Document()
    set_page(doc, (14, 15, 13, 15))
    configure_styles(doc, 8.4, "075E50")
    pattern = tmp / "stilvoll-pattern.png"
    photo = tmp / "stilvoll-photo.png"
    draw_stilvoll_pattern(pattern)
    create_placeholder_photo(photo, ("36B873", "D9F2E5"))
    add_header_background(doc, pattern)
    header = doc.add_table(rows=1, cols=2)
    set_table_geometry(header, [8640, 1584])
    remove_table_borders(header)
    identity, photo_cell = header.rows[0].cells
    add_text(identity, "{{VORNAME}} {{NACHNAME}}", size=23, color="075E50", after=2)
    add_text(identity, "{{BERUFSBEZEICHNUNG}} | {{FACHGEBIETE}}", size=12, color="36B873", after=4)
    add_text(
        identity,
        "{{TELEFON}}  ·  {{EMAIL}}  ·  {{LINKEDIN}}  ·  {{ORT}}  ·  {{GEBURTSZEILE}}",
        size=7.8,
        color=INK,
        line=1.1,
    )
    photo_paragraph = photo_cell.paragraphs[0]
    photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    shape = photo_paragraph.add_run().add_picture(str(photo), width=Mm(26), height=Mm(26))
    set_photo_alt_text(shape)
    body = doc.add_table(rows=1, cols=3)
    set_table_geometry(body, [3061, 624, 6519])
    remove_table_borders(body)
    left, _, right = body.rows[0].cells
    add_section_title(left, "{{ZUSAMMENFASSUNG_TITEL}}", color=MUTED, size=9.5, border_color="AEB8B5")
    add_text(left, "{{ZUSAMMENFASSUNG}}", size=8.4, line=1.25)
    add_section_title(left, "{{STAERKEN_TITEL}}", color=MUTED, size=9.5, border_color="AEB8B5")
    for index in range(1, 5):
        add_text(left, f"●  {{{{STAERKE_{index}_TITEL}}}}", size=9.2, color="075E50", bold=True, after=1)
        add_text(left, f"{{{{STAERKE_{index}_BESCHREIBUNG}}}}", size=8.1, after=3, line=1.2)
    add_section_title(left, "{{SPRACHEN_TITEL}}", color=MUTED, size=9.5, border_color="AEB8B5")
    for index in range(1, 4):
        add_text(left, f"{{{{SPRACHE_{index}}}}}  {{{{SPRACHNIVEAU_{index}}}}}", size=8.2, color="075E50", after=2)
    add_section_title(right, "{{ERFAHRUNG_TITEL}}", color=MUTED, size=9.5, border_color="AEB8B5")
    for index in range(1, 5):
        add_career_entry(right, index, primary="075E50", accent="36B873")
    add_section_title(right, "{{AUSBILDUNG_TITEL}}", color=MUTED, size=9.5, border_color="AEB8B5")
    for index in range(1, 3):
        add_text(right, f"{{{{ABSCHLUSS_{index}}}}}", size=10.5, color="075E50", after=1)
        add_text(right, f"{{{{HOCHSCHULE_{index}}}}}", size=9.3, color="36B873", after=1)
        add_text(
            right,
            f"{{{{AUSBILDUNG_START_{index}}}}}–{{{{AUSBILDUNG_ENDE_{index}}}}} · {{{{AUSBILDUNG_ORT_{index}}}}}",
            size=7.8,
            color=MUTED,
            after=3,
        )
    add_footer(doc, "075E50")
    doc.save(output)


def build_einfach_visual(output: Path, tmp: Path) -> None:
    doc = Document()
    set_page(doc, (14, 15, 12, 15))
    configure_styles(doc, 8.5, "073B8F")
    pattern = tmp / "einfach-pattern.png"
    photo = tmp / "einfach-photo.png"
    draw_einfach_pattern(pattern)
    create_placeholder_photo(photo, ("4AA7F5", "EAF6FD"))
    add_header_background(doc, pattern)
    header = doc.add_table(rows=1, cols=2)
    set_table_geometry(header, [8410, 1814])
    remove_table_borders(header)
    identity, photo_cell = header.rows[0].cells
    add_text(identity, "{{VORNAME}} {{NACHNAME}}", size=24, color="073B8F", bold=True, after=2)
    add_text(identity, "{{BERUFSBEZEICHNUNG}} | {{FACHGEBIETE}}", size=11.5, color="4AA7F5", bold=True, after=3)
    add_text(
        identity,
        "{{TELEFON}}                {{EMAIL}}\n{{LINKEDIN}}                {{ORT}}\n{{GEBURTSZEILE}}",
        size=7.8,
        color=INK,
        bold=True,
        line=1.15,
    )
    photo_paragraph = photo_cell.paragraphs[0]
    photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    shape = photo_paragraph.add_run().add_picture(str(photo), width=Mm(30), height=Mm(30))
    set_photo_alt_text(shape)
    add_section_title(doc, "{{ZUSAMMENFASSUNG_TITEL}}", color="073B8F", size=14, border_color="073B8F")
    add_text(doc, "{{ZUSAMMENFASSUNG}}", size=8.5, line=1.25)
    add_section_title(doc, "{{STAERKEN_TITEL}}", color="073B8F", size=14, border_color="073B8F")
    strengths = doc.add_table(rows=2, cols=2)
    set_table_geometry(strengths, [5112, 5112])
    remove_table_borders(strengths)
    for offset, cell in enumerate([cell for row in strengths.rows for cell in row.cells], start=1):
        add_text(cell, f"✣  {{{{STAERKE_{offset}_TITEL}}}}", size=9.5, color="073B8F", bold=True, after=1)
        add_text(cell, f"{{{{STAERKE_{offset}_BESCHREIBUNG}}}}", size=8.2, line=1.2)
    add_section_title(doc, "{{ERFAHRUNG_TITEL}}", color="073B8F", size=14, border_color="073B8F")
    for index in range(1, 5):
        add_career_entry(doc, index, primary="073B8F", accent="4AA7F5")
    add_section_title(doc, "{{AUSBILDUNG_TITEL}}", color="073B8F", size=14, border_color="073B8F")
    for index in range(1, 3):
        add_text(doc, f"{{{{ABSCHLUSS_{index}}}}}", size=10.5, color="073B8F", after=1)
        add_text(doc, f"{{{{HOCHSCHULE_{index}}}}}", size=9.3, color="4AA7F5", bold=True, after=1)
        add_text(
            doc,
            f"{{{{AUSBILDUNG_START_{index}}}}}–{{{{AUSBILDUNG_ENDE_{index}}}}} · {{{{AUSBILDUNG_ORT_{index}}}}}",
            size=7.8,
            color=MUTED,
            after=3,
        )
    add_section_title(doc, "{{SPRACHEN_TITEL}}", color="073B8F", size=14, border_color="073B8F")
    language_table = doc.add_table(rows=1, cols=2)
    set_table_geometry(language_table, [5112, 5112])
    remove_table_borders(language_table)
    for index, cell in enumerate(language_table.rows[0].cells, start=1):
        add_text(cell, f"{{{{SPRACHE_{index}}}}}", size=9, color="073B8F", bold=True, after=1)
        add_text(cell, f"{{{{SPRACHNIVEAU_{index}}}}}  ● ● ● ● ○", size=8.2)
    add_footer(doc, "073B8F")
    doc.save(output)


def build_ats(output: Path, primary: str) -> None:
    doc = Document()
    set_page(doc, (14, 15, 14, 15))
    configure_styles(doc, 9, primary)
    add_text(doc, "{{VORNAME}} {{NACHNAME}}", size=20, color=primary, bold=True, after=2)
    add_text(doc, "{{BERUFSBEZEICHNUNG}}", size=10.5, color=primary, bold=True, after=2)
    add_text(
        doc,
        "{{TELEFON}} · {{EMAIL}} · {{LINKEDIN}} · {{WEBSITE}} · {{ORT}} · {{GEBURTSZEILE}}",
        size=8.2,
        color=INK,
        after=6,
    )
    add_section_title(doc, "Zusammenfassung", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    add_text(doc, "{{ZUSAMMENFASSUNG}}", size=9, line=1.2)
    add_section_title(doc, "Berufserfahrung", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    for index in range(1, 7):
        add_career_entry(doc, index, primary=primary, accent=primary)
    add_section_title(doc, "Ausbildung", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    for index in range(1, 4):
        add_text(doc, f"{{{{ABSCHLUSS_{index}}}}}", size=10, color=primary, bold=True, after=1)
        add_text(doc, f"{{{{HOCHSCHULE_{index}}}}} · {{{{AUSBILDUNG_START_{index}}}}}–{{{{AUSBILDUNG_ENDE_{index}}}}} · {{{{AUSBILDUNG_ORT_{index}}}}}", size=8.5, after=2)
    add_section_title(doc, "Kenntnisse", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    add_text(doc, "{{KENNTNISSE}}", size=9)
    add_section_title(doc, "Sprachen", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    add_text(doc, "{{SPRACHEN_ATS}}", size=9)
    add_section_title(doc, "Stärken", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    add_text(doc, "{{STAERKEN_ATS}}", size=9)
    add_section_title(doc, "Zertifikate", color=primary, size=11, border_color="AEB8BF", uppercase=False)
    add_text(doc, "{{ZERTIFIKATE}}", size=9)
    add_footer(doc, primary)
    doc.save(output)


def load_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textlength(candidate, font=font) <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def preview_stilvoll(path: Path) -> None:
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    dark, green, muted, line = "#075E50", "#36B873", "#6D777C", "#EEF1F0"
    for x in range(400, 1220, 210):
        for y in range(-40, 650, 140):
            draw.line((x - 100, y + 100, x, y, x + 100, y + 100), fill=line, width=2)
            draw.line((x - 100, y + 122, x, y + 22, x + 100, y + 122), fill=line, width=2)
    draw.text((88, 82), "SEBASTIAN WAGNER", font=load_font(50), fill=dark)
    draw.text((88, 142), "Pflegefachmann | Notfallmedizin | Patientenbetreuung", font=load_font(26), fill=green)
    draw.text((88, 194), "+49 30 12345678   ·   sebastian@example.de   ·   linkedin.com   ·   Berlin", font=load_font(17), fill="#465156")
    draw.rounded_rectangle((1010, 82, 1145, 217), 8, fill="#D9F2E5")
    draw.ellipse((1042, 103, 1113, 174), fill=green)
    draw.rounded_rectangle((1025, 160, 1130, 235), 25, fill=green)
    left_x, right_x, top = 88, 470, 288
    def section(x: int, y: int, width: int, title: str):
        draw.text((x, y), title, font=load_font(17), fill=muted)
        draw.line((x, y + 27, x + width, y + 27), fill="#AEB8B5", width=2)
        return y + 52
    y = section(left_x, top, 315, "ZUSAMMENFASSUNG")
    body = "Mit über 10 Jahren Erfahrung in der Notfallmedizin verbinde ich exzellente Patientenversorgung mit klarer Teamführung."
    for line_text in wrap(draw, body, load_font(17), 300):
        draw.text((left_x, y), line_text, font=load_font(17), fill="#465156"); y += 25
    y += 24
    y = section(left_x, y, 315, "STÄRKEN")
    for title, text in [("Patientenfokus", "Stets im Interesse der Patienten handelnd."), ("Teamarbeit", "Vertrauensvolle Zusammenarbeit in Teams."), ("Weiterentwicklung", "Aktive Optimierung von Arbeitsstrukturen.")]:
        draw.ellipse((left_x, y, left_x + 48, y + 48), fill="#F1F3F2")
        draw.ellipse((left_x + 18, y + 18, left_x + 30, y + 30), fill=green)
        draw.text((left_x + 65, y), title, font=load_font(20), fill=dark)
        for line_text in wrap(draw, text, load_font(16), 235):
            draw.text((left_x + 65, y + 28), line_text, font=load_font(16), fill="#465156"); y += 20
        y += 46
    y = section(left_x, max(y + 35, 1130), 315, "SPRACHEN")
    draw.text((left_x, y), "Deutsch    Muttersprache", font=load_font(17), fill=dark)
    draw.text((left_x, y + 55), "Englisch   Erweitert", font=load_font(17), fill=dark)
    for row, filled in ((y + 9, 5), (y + 64, 3)):
        for index in range(5):
            draw.ellipse((335 + index * 13, row, 343 + index * 13, row + 8), fill=dark if index < filled else "#DDE2E0")
    y = section(right_x, top, 675, "ERFAHRUNG")
    experiences = [
        ("Pflegefachmann Notfallmedizin", "Universitätsklinikum Düsseldorf", "2017–Heute · Düsseldorf"),
        ("Senior Pflegefachkraft", "Klinikum Stuttgart", "2014–2017 · Stuttgart"),
        ("Pflegespezialist Notaufnahme", "Universitätsklinikum Hamburg", "2011–2014 · Hamburg"),
    ]
    for role, company, meta in experiences:
        draw.text((right_x, y), role, font=load_font(23), fill=dark)
        draw.text((right_x, y + 34), company, font=load_font(20), fill=green)
        draw.text((920, y + 36), meta, font=load_font(15), fill=muted)
        y += 70
        for bullet in ["Optimierte Abläufe und steigerte die Effizienz messbar.", "Führte interdisziplinäre Teams und Schulungsprogramme.", "Verbesserte die Qualität der Patientenbetreuung."]:
            draw.text((right_x + 8, y), "·", font=load_font(17), fill=muted)
            draw.text((right_x + 28, y), bullet, font=load_font(16), fill="#465156")
            y += 28
        y += 35
    y = section(right_x, y + 30, 675, "AUSBILDUNG")
    draw.text((right_x, y), "Master in Klinischer Pflege", font=load_font(22), fill=dark)
    draw.text((right_x, y + 34), "Hochschule für Gesundheit Bochum", font=load_font(19), fill=green)
    draw.text((930, y + 36), "2009–2011 · Bochum", font=load_font(15), fill=muted)
    draw.text((88, 1690), "portfolio.example.com", font=load_font(14), fill=dark)
    draw.text((1080, 1690), "Seite 1 / 1", font=load_font(14), fill=muted)
    image.save(path)


def preview_einfach(path: Path) -> None:
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    blue, light, text, pattern = "#073B8F", "#4AA7F5", "#424B50", "#EAF6FD"
    for points in [[(750, -20), (750, 80), (840, 80), (840, -20)], [(280, 150), (280, 270), (370, 270), (370, 200)], [(910, 120), (910, 250), (1010, 250), (1010, 190)], [(680, 1530), (680, 1700), (770, 1700)]]:
        draw.line(points, fill=pattern, width=28, joint="curve")
    draw.text((88, 82), "ANNA KOCH", font=load_font(50, True), fill=blue)
    draw.text((88, 142), "Regionaler Vertriebsleiter | Baustoffe | Teamführung", font=load_font(25, True), fill=light)
    draw.text((88, 190), "• +49 30 12345678                      • anna@example.de", font=load_font(17, True), fill=text)
    draw.text((88, 220), "• linkedin.com                           • Berlin, Deutschland", font=load_font(17, True), fill=text)
    draw.ellipse((1000, 72, 1150, 222), fill="#D8EBFA")
    draw.ellipse((1040, 102, 1110, 172), fill=light)
    draw.rounded_rectangle((1019, 154, 1131, 235), 28, fill=light)
    def section(y: int, title: str):
        draw.text((88, y), title, font=load_font(29, True), fill=blue)
        draw.line((88, y + 42, 1150, y + 42), fill=blue, width=4)
        return y + 65
    y = section(292, "ZUSAMMENFASSUNG")
    summary = "Mit über 13 Jahren Erfahrung im Vertrieb, einer starken Führungsbilanz und erfolgreichen Produkteinführungen positioniere ich mich ideal für diese Rolle."
    for line_text in wrap(draw, summary, load_font(17), 1060):
        draw.text((88, y), line_text, font=load_font(17), fill=text); y += 25
    y = section(468, "STÄRKEN")
    draw.ellipse((100, y + 5, 116, y + 21), fill=light)
    draw.text((145, y), "Führungskompetenz", font=load_font(19, True), fill=blue)
    draw.text((145, y + 34), "Erfolgreiche Leitung eines Teams und messbare Umsatzsteigerung.", font=load_font(16), fill=text)
    draw.ellipse((654, y + 5, 670, y + 21), fill=light)
    draw.text((700, y), "Vertriebsstrategie", font=load_font(19, True), fill=blue)
    draw.text((700, y + 34), "Entwicklung und Implementierung einer erfolgreichen Strategie.", font=load_font(16), fill=text)
    y = section(660, "ERFAHRUNG")
    for role, company, meta, bullets in [
        ("Bereichsleiter Vertrieb", "BASF", "2018–2023 · Ludwigshafen", 4),
        ("Vertriebsleiter", "HeidelbergCement", "2014–2018 · Heidelberg", 3),
        ("Key Account Manager", "Saint-Gobain", "2010–2014 · Düsseldorf", 3),
    ]:
        draw.text((88, y), role, font=load_font(22), fill=blue)
        draw.text((88, y + 32), company, font=load_font(18, True), fill=light)
        draw.text((88, y + 62), meta, font=load_font(15), fill="#667177")
        y += 92
        for index in range(bullets):
            draw.text((90, y), "•", font=load_font(15), fill=text)
            draw.text((115, y), f"Messbarer Erfolg in Vertrieb und Teamführung {index + 1}.", font=load_font(15), fill=text)
            y += 24
        draw.line((88, y + 8, 1150, y + 8), fill="#D4D9DC", width=1)
        y += 30
    y = section(1370, "AUSBILDUNG")
    draw.text((88, y), "Master in Betriebswirtschaft", font=load_font(21), fill=blue)
    draw.text((88, y + 33), "Universität München", font=load_font(18, True), fill=light)
    draw.text((88, y + 62), "2008–2010 · München, Deutschland", font=load_font(15), fill="#667177")
    y = section(1518, "SPRACHEN")
    draw.text((88, y), "Deutsch   Muttersprache     ● ● ● ● ●", font=load_font(17, True), fill=blue)
    draw.text((570, y), "Englisch   Erweitert        ● ● ● ○ ○", font=load_font(17, True), fill=blue)
    draw.text((88, 1690), "portfolio.example.com", font=load_font(14), fill=blue)
    draw.text((1080, 1690), "Seite 1 / 1", font=load_font(14), fill="#667177")
    image.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path("public/templates"))
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    tmp = output_dir / ".managed-resume-assets"
    tmp.mkdir(exist_ok=True)
    outputs = {
        "stilvoll_visual": output_dir / "Stilvoll_Lebenslauf_Muster.docx",
        "stilvoll_ats": output_dir / "Stilvoll_Lebenslauf_ATS.docx",
        "stilvoll_preview": output_dir / "Stilvoll_Lebenslauf_Muster.preview.png",
        "einfach_visual": output_dir / "Einfach_Lebenslauf_Muster.docx",
        "einfach_ats": output_dir / "Einfach_Lebenslauf_ATS.docx",
        "einfach_preview": output_dir / "Einfach_Lebenslauf_Muster.preview.png",
    }
    if not args.force:
        existing = [path for path in outputs.values() if path.exists()]
        if existing:
            raise SystemExit(f"Refusing to overwrite: {existing[0]}")
    build_stilvoll_visual(outputs["stilvoll_visual"], tmp)
    build_ats(outputs["stilvoll_ats"], "075E50")
    preview_stilvoll(outputs["stilvoll_preview"])
    build_einfach_visual(outputs["einfach_visual"], tmp)
    build_ats(outputs["einfach_ats"], "073B8F")
    preview_einfach(outputs["einfach_preview"])
    for child in tmp.iterdir():
        child.unlink()
    tmp.rmdir()
    for path in outputs.values():
        print(path)


if __name__ == "__main__":
    main()
