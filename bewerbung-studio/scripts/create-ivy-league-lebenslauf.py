#!/usr/bin/env python3
"""Create the bundled Ivy League Word resume templates and preview.

The visual DOCX is a classic A4 single-column resume with an original pastel
watercolor background. The ATS DOCX uses a plain linear document structure.
Existing outputs are only replaced when --force is supplied.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips


# compact_reference_guide with named "Ivy League A4 resume" overrides.
# First-page treatment: proposal_centerpiece adapted to a compact CV header.
PAGE_WIDTH_DXA = 11_906
PAGE_HEIGHT_DXA = 16_838
MARGIN_X_DXA = 794  # 14 mm
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - (2 * MARGIN_X_DXA)
STRENGTH_COLUMN_DXA = CONTENT_WIDTH_DXA // 3
PRIMARY = "073C8C"
ACCENT = "FF6A00"
INK = "3F4B50"
MUTED = "667177"
LINE = "0B459A"
INACTIVE = "DCE9E8"
BODY_FONT = "Arial"
HEADING_FONT = "Georgia"


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    size: float = 8.5,
    color: str = INK,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def configure_style(
    doc: Document,
    name: str,
    *,
    font: str = BODY_FONT,
    size: float,
    color: str,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.0,
    alignment=None,
) -> None:
    style = (
        doc.styles[name]
        if name in doc.styles
        else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    style.font.name = font
    r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line
    if alignment is not None:
        style.paragraph_format.alignment = alignment


def configure_styles(doc: Document, *, ats: bool = False) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial" if ats else BODY_FONT
    r_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Arial" if ats else BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), "Arial" if ats else BODY_FONT)
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.12 if ats else 1.06

    configure_style(
        doc,
        "CV Name",
        font="Arial" if ats else HEADING_FONT,
        size=19 if ats else 17.5,
        color=PRIMARY,
        bold=True,
        after=2,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    configure_style(
        doc,
        "CV Title",
        size=10 if ats else 11.5,
        color=PRIMARY if ats else ACCENT,
        after=2,
        line=1.08,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    configure_style(
        doc,
        "Contact",
        size=7.8,
        color=INK,
        after=1,
        line=1.05,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    configure_style(
        doc,
        "Section",
        font="Arial" if ats else HEADING_FONT,
        size=11 if ats else 13.5,
        color=PRIMARY,
        bold=True,
        before=7 if ats else 8,
        after=4,
        alignment=WD_ALIGN_PARAGRAPH.LEFT
        if ats
        else WD_ALIGN_PARAGRAPH.CENTER,
    )
    configure_style(
        doc,
        "Summary",
        size=8.5,
        color=INK,
        after=1,
        line=1.15,
    )
    configure_style(
        doc,
        "Strength Title",
        size=9.5,
        color=PRIMARY,
        bold=True,
        after=1,
        line=1.05,
    )
    configure_style(
        doc,
        "Strength Description",
        size=8.1,
        color=INK,
        after=0,
        line=1.08,
    )
    configure_style(
        doc,
        "Entry Organization",
        size=10.5,
        color=ACCENT,
        after=1,
        line=1.05,
    )
    configure_style(
        doc,
        "Entry Role",
        size=9.7,
        color=PRIMARY,
        after=1,
        line=1.05,
    )
    configure_style(
        doc,
        "Entry Bullet",
        size=8.3,
        color=INK,
        after=1,
        line=1.08,
    )
    configure_style(
        doc,
        "Optional",
        size=8.3,
        color=INK,
        after=1,
        line=1.08,
    )
    configure_style(
        doc,
        "Cell Terminator",
        size=1,
        color="FFFFFF",
        after=0,
        line=1.0,
    )

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(8.3)
    bullet.font.color.rgb = RGBColor.from_string(INK)
    bullet.paragraph_format.left_indent = Mm(4.5)
    bullet.paragraph_format.first_line_indent = Mm(-2.2)
    bullet.paragraph_format.space_after = Pt(1)
    bullet.paragraph_format.line_spacing = 1.08


def configure_page(doc: Document, *, ats: bool = False) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(13 if not ats else 14)
    section.right_margin = Mm(14)
    section.bottom_margin = Mm(12)
    section.left_margin = Mm(14)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(6)


def paragraph_bottom_border(paragraph, color: str = LINE) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_margins(
    cell,
    *,
    top: int = 45,
    bottom: int = 45,
    start: int = 70,
    end: int = 70,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "nil")
        borders.append(border)


def apply_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "0")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        row.height = None
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def remove_initial_cell_paragraph(cell) -> None:
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        paragraph = cell.paragraphs[0]
        paragraph._element.getparent().remove(paragraph._element)


def add_section_heading(container, text: str):
    paragraph = container.add_paragraph(text, style="Section")
    paragraph_bottom_border(paragraph)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_tabbed_paragraph(
    container,
    left: str,
    right: str,
    *,
    style: str,
    right_color: str = INK,
    keep_with_next: bool = True,
):
    paragraph = container.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Mm(182), WD_TAB_ALIGNMENT.RIGHT
    )
    left_run = paragraph.add_run(left)
    left_style = container.part.document.styles[style]
    set_run_font(
        left_run,
        size=left_style.font.size.pt,
        color=str(left_style.font.color.rgb),
        name=left_style.font.name,
        bold=bool(left_style.font.bold),
    )
    paragraph.add_run("\t")
    right_run = paragraph.add_run(right)
    set_run_font(right_run, size=8.3, color=right_color)
    return paragraph


def create_watercolor(path: Path) -> None:
    width, height = 1654, 2339
    base = Image.new("RGB", (width, height), "#F8FBF8")
    layer = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(layer)
    blobs = [
        ((-240, -120, 720, 750), (206, 235, 241, 135)),
        ((780, -180, 1870, 650), (255, 237, 204, 120)),
        ((-180, 620, 800, 1450), (215, 240, 230, 100)),
        ((900, 520, 1870, 1460), (216, 237, 241, 108)),
        ((-260, 1460, 850, 2440), (221, 242, 239, 110)),
        ((580, 1360, 1790, 2420), (255, 235, 192, 115)),
        ((1040, 1900, 1850, 2500), (204, 235, 239, 128)),
    ]
    for box, fill in blobs:
        draw.ellipse(box, fill=fill)
    layer = layer.filter(ImageFilter.GaussianBlur(105))
    base = Image.alpha_composite(base.convert("RGBA"), layer)
    grain = Image.effect_noise((width, height), 18).convert("L")
    grain_alpha = grain.point(lambda value: int(value * 0.045))
    grain_layer = Image.new("RGBA", (width, height), (231, 240, 236, 0))
    grain_layer.putalpha(grain_alpha)
    base = Image.alpha_composite(base, grain_layer)
    base.convert("RGB").save(path, "PNG", optimize=True)


def convert_inline_to_page_background(shape) -> None:
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "0",
        "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for axis in ("H", "V"):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    for element in (
        inline.find(qn("wp:extent")),
        inline.find(qn("wp:effectExtent")),
    ):
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element in (
        inline.find(qn("wp:docPr")),
        inline.find(qn("wp:cNvGraphicFramePr")),
        inline.find(qn("a:graphic")),
    ):
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def add_watercolor_background(doc: Document, image_path: Path) -> None:
    paragraph = doc.sections[0].header.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Mm(210), height=Mm(297))
    shape._inline.docPr.set("descr", "Originaler Pastell-Suluboya-Hintergrund")
    convert_inline_to_page_background(shape)


def add_page_field(paragraph, field_name: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_name} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    table = footer.add_table(rows=1, cols=2, width=Mm(182))
    set_table_borders_none(table)
    apply_table_geometry(table, [CONTENT_WIDTH_DXA // 2, CONTENT_WIDTH_DXA // 2])
    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
    left_p = left.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = left_p.add_run("{{WEBSITE}}")
    set_run_font(run, size=7.1, color=PRIMARY)
    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right_p.add_run("Seite ")
    set_run_font(run, size=7.1, color=MUTED)
    add_page_field(right_p, "PAGE")
    run = right_p.add_run(" / ")
    set_run_font(run, size=7.1, color=MUTED)
    add_page_field(right_p, "NUMPAGES")


def add_experience(container, index: int, *, ats: bool = False) -> None:
    organization = f"{{{{UNTERNEHMEN_{index}}}}}"
    location = f"{{{{ARBEITSORT_{index}}}}}"
    role = f"{{{{POSITION_{index}}}}}"
    date = (
        f"{{{{STARTDATUM_{index}}}}}{{{{DATUM_TRENNER_{index}}}}}"
        f"{{{{ENDDATUM_{index}}}}}"
    )
    add_tabbed_paragraph(
        container,
        organization,
        location,
        style="Entry Organization",
    )
    add_tabbed_paragraph(container, role, date, style="Entry Role")
    container.add_paragraph(f"{{{{BESCHREIBUNG_{index}}}}}", style="Optional")
    for achievement_index in range(1, 6):
        paragraph = container.add_paragraph(
            f"{{{{ERFOLG_{index}_{achievement_index}}}}}",
            style="List Bullet",
        )
        paragraph.paragraph_format.keep_together = True
    if ats:
        container.add_paragraph(
            f"{{{{TECHNOLOGIEN_{index}}}}}", style="Optional"
        )


def add_education(container, index: int) -> None:
    add_tabbed_paragraph(
        container,
        f"{{{{HOCHSCHULE_{index}}}}}",
        f"{{{{AUSBILDUNG_ORT_{index}}}}}",
        style="Entry Organization",
    )
    add_tabbed_paragraph(
        container,
        f"{{{{ABSCHLUSS_{index}}}}} {{{{FACHRICHTUNG_{index}}}}}",
        (
            f"{{{{AUSBILDUNG_START_{index}}}}}"
            f"{{{{AUSBILDUNG_DATUM_TRENNER_{index}}}}}"
            f"{{{{AUSBILDUNG_ENDE_{index}}}}}"
        ),
        style="Entry Role",
    )


def build_visual(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    watercolor_path = output_path.parent / ".ivy-league-watercolor.png"
    create_watercolor(watercolor_path)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    add_watercolor_background(doc, watercolor_path)
    add_footer(doc)

    doc.add_paragraph("{{VORNAME}} {{NACHNAME}}", style="CV Name")
    doc.add_paragraph(
        "{{BERUFSBEZEICHNUNG}} | {{FACHGEBIETE}}", style="CV Title"
    )
    doc.add_paragraph(
        "{{KONTAKT_ZEILE_1}} · {{KONTAKT_ZEILE_2}} · {{KONTAKT_ZEILE_3}}",
        style="Contact",
    )
    doc.add_paragraph(
        "{{HEADER_KONTAKT_4}} · {{HEADER_KONTAKT_5}} · {{HEADER_KONTAKT_6}}",
        style="Contact",
    )

    add_section_heading(doc, "{{ZUSAMMENFASSUNG_TITEL}}")
    doc.add_paragraph("{{ZUSAMMENFASSUNG}}", style="Summary")

    add_section_heading(doc, "{{STAERKEN_TITEL}}")
    strengths = doc.add_table(rows=1, cols=3)
    set_table_borders_none(strengths)
    apply_table_geometry(
        strengths,
        [
            STRENGTH_COLUMN_DXA,
            STRENGTH_COLUMN_DXA,
            CONTENT_WIDTH_DXA - (2 * STRENGTH_COLUMN_DXA),
        ],
    )
    for index, cell in enumerate(strengths.rows[0].cells, start=1):
        remove_initial_cell_paragraph(cell)
        title = cell.add_paragraph(style="Strength Title")
        icon = title.add_run(("♥", "●", "●")[index - 1] + " ")
        set_run_font(icon, size=12, color=ACCENT)
        value = title.add_run(f"{{{{STAERKE_{index}_TITEL}}}}")
        set_run_font(value, size=9.5, color=PRIMARY, bold=True)
        cell.add_paragraph(
            f"{{{{STAERKE_{index}_BESCHREIBUNG}}}}",
            style="Strength Description",
        )
        cell.add_paragraph(" ", style="Cell Terminator")

    add_section_heading(doc, "{{ERFAHRUNG_TITEL}}")
    for index in range(1, 7):
        add_experience(doc, index)

    add_section_heading(doc, "{{AUSBILDUNG_TITEL}}")
    for index in range(1, 4):
        add_education(doc, index)

    add_section_heading(doc, "{{KENNTNISSE_TITEL}}")
    doc.add_paragraph("{{KENNTNISSE}}", style="Optional")

    add_section_heading(doc, "{{SPRACHEN_TITEL}}")
    languages = doc.add_table(rows=2, cols=2)
    set_table_borders_none(languages)
    apply_table_geometry(
        languages,
        [CONTENT_WIDTH_DXA // 2, CONTENT_WIDTH_DXA - CONTENT_WIDTH_DXA // 2],
    )
    for index, cell in enumerate(
        [cell for row in languages.rows for cell in row.cells],
        start=1,
    ):
        remove_initial_cell_paragraph(cell)
        if index <= 3:
            name = cell.add_paragraph(style="Entry Role")
            name.add_run(f"{{{{SPRACHE_{index}}}}}")
            level = cell.add_paragraph(style="Optional")
            level.add_run(
                f"{{{{SPRACHNIVEAU_{index}}}}}  {{{{SPRACHE_{index}_PUNKTE}}}}"
            )
        cell.add_paragraph(" ", style="Cell Terminator")

    for title_key, content_key in (
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
        ("PROJEKTE_TITEL", "PROJEKTE"),
        ("WEITERBILDUNGEN_TITEL", "WEITERBILDUNGEN"),
        ("VEROEFFENTLICHUNGEN_TITEL", "VEROEFFENTLICHUNGEN"),
        ("EHRENAMT_TITEL", "EHRENAMT"),
    ):
        add_section_heading(doc, f"{{{{{title_key}}}}}")
        doc.add_paragraph(f"{{{{{content_key}}}}}", style="Optional")

    doc.save(output_path)
    watercolor_path.unlink(missing_ok=True)
    return output_path


def build_ats(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc, ats=True)
    configure_styles(doc, ats=True)
    add_footer(doc)

    doc.add_paragraph("{{VORNAME}} {{NACHNAME}}", style="CV Name")
    doc.add_paragraph("{{BERUFSBEZEICHNUNG}}", style="CV Title")
    doc.add_paragraph("{{FACHGEBIETE}}", style="Contact")
    doc.add_paragraph("{{KONTAKT_ZEILE_1}}", style="Contact")
    doc.add_paragraph("{{KONTAKT_ZEILE_2}}", style="Contact")
    doc.add_paragraph("{{KONTAKT_ZEILE_3}}", style="Contact")

    add_section_heading(doc, "{{ZUSAMMENFASSUNG_TITEL}}")
    doc.add_paragraph("{{ZUSAMMENFASSUNG}}", style="Summary")
    add_section_heading(doc, "{{BERUFSERFAHRUNG_TITEL}}")
    for index in range(1, 7):
        add_experience(doc, index, ats=True)
    add_section_heading(doc, "{{AUSBILDUNG_TITEL}}")
    for index in range(1, 4):
        add_education(doc, index)
    for title_key, content_key in (
        ("KENNTNISSE_TITEL", "KENNTNISSE"),
        ("SPRACHEN_TITEL", "SPRACHEN_ATS"),
        ("STAERKEN_TITEL", "STAERKEN_ATS"),
        ("PROJEKTE_TITEL", "PROJEKTE"),
        ("ZERTIFIKATE_TITEL", "ZERTIFIKATE"),
        ("WEITERBILDUNGEN_TITEL", "WEITERBILDUNGEN"),
    ):
        add_section_heading(doc, f"{{{{{title_key}}}}}")
        doc.add_paragraph(f"{{{{{content_key}}}}}", style="Optional")
    doc.save(output_path)
    return output_path


def font(size: int, *, bold: bool = False, serif: bool = False):
    candidates = (
        [
            "C:/Windows/Fonts/georgiab.ttf"
            if bold
            else "C:/Windows/Fonts/georgia.ttf"
        ]
        if serif
        else [
            "C:/Windows/Fonts/arialbd.ttf"
            if bold
            else "C:/Windows/Fonts/arial.ttf"
        ]
    )
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def build_preview(output_path: Path, *, force: bool) -> Path:
    if output_path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing file: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    create_watercolor(output_path)
    image = Image.open(output_path).convert("RGB").resize((900, 1273))
    draw = ImageDraw.Draw(image)
    blue = "#073C8C"
    orange = "#FF6A00"
    ink = "#3F4B50"
    muted = "#667177"
    left, right = 60, 840

    def center(text: str, y: int, text_font, fill: str) -> int:
        box = draw.textbbox((0, 0), text, font=text_font)
        draw.text(((900 - (box[2] - box[0])) / 2, y), text, font=text_font, fill=fill)
        return y + box[3] - box[1]

    def heading(text: str, y: int) -> int:
        center(text, y, font(24, bold=True, serif=True), blue)
        draw.line((left, y + 31, right, y + 31), fill=blue, width=2)
        return y + 43

    def wrap(text: str, max_width: int, text_font):
        words = text.split()
        lines: list[str] = []
        line = ""
        for word in words:
            candidate = f"{line} {word}".strip()
            if draw.textlength(candidate, font=text_font) <= max_width:
                line = candidate
            else:
                if line:
                    lines.append(line)
                line = word
        if line:
            lines.append(line)
        return lines

    y = 55
    y = center("LENA HOFFMANN", y, font(27, bold=True, serif=True), blue) + 6
    y = center(
        "Ingenieurin | Maschinenbau | Regelwerke | FEM",
        y,
        font(16),
        orange,
    ) + 5
    y = center(
        "+49 30 12345678  •  lena@example.de  •  linkedin.com  •  München",
        y,
        font(10),
        ink,
    ) + 22

    y = heading("Zusammenfassung", y)
    summary = (
        "Engagierte Ingenieurin mit über acht Jahren Erfahrung im Maschinenbau "
        "und in der Entwurfsprüfung. Verbindet technische Präzision mit klarer "
        "Projektführung und messbaren Verbesserungen."
    )
    for line in wrap(summary, right - left, font(11)):
        draw.text((left, y), line, font=font(11), fill=ink)
        y += 17
    y += 15

    y = heading("Stärken", y)
    strengths = [
        ("♥", "Analysefähigkeit", "Präzise Bewertung komplexer Systeme"),
        ("●", "Kundenbetreuung", "Verlässliche Beratung und Schulung"),
        ("●", "Teamführung", "Führung interdisziplinärer Teams"),
    ]
    column_width = 250
    for index, (icon, title, description) in enumerate(strengths):
        x = left + index * 265
        draw.text((x, y), icon, font=font(20), fill=orange)
        draw.text((x + 24, y), title, font=font(12, bold=True), fill=blue)
        line_y = y + 20
        for line in wrap(description, column_width - 24, font(10)):
            draw.text((x + 24, line_y), line, font=font(10), fill=ink)
            line_y += 15
    y += 68

    y = heading("Erfahrung", y)
    entries = [
        (
            "Siemens AG",
            "Berlin",
            "Senior Maschinenbauingenieurin",
            "2019 – 2023",
            [
                "Leitung eines Teams von fünf Ingenieurinnen und Ingenieuren.",
                "Projekteffizienz um 15 % und Durchlaufzeit um 25 % verbessert.",
                "Neue Analysemethode für sichere Druckbehälter eingeführt.",
                "Fertigungsprozesse optimiert und Kosten deutlich reduziert.",
                "Markteinführung durch klare Projektsteuerung beschleunigt.",
            ],
        ),
        (
            "Bosch Rexroth",
            "Stuttgart",
            "Prüfingenieurin",
            "2016 – 2019",
            [
                "Berechnungsgenauigkeit mit ROHR2 um 20 % gesteigert.",
                "Mehr als 50 Kundenprojekte erfolgreich begleitet.",
                "Schulungsprogramm für neue Teammitglieder entwickelt.",
                "Werkstoffeinsatz und Prüfverfahren nachhaltig optimiert.",
            ],
        ),
        (
            "BASF SE",
            "Ludwigshafen",
            "Fachingenieurin Verfahrenstechnik",
            "2013 – 2016",
            [
                "Technische Regelwerke digitalisiert und Wissenstransfer beschleunigt.",
                "Mehr als 100 interne Richtlinien strukturiert dokumentiert.",
                "200 Mitarbeitende zu technischen Standards geschult.",
            ],
        ),
    ]
    for company, city, role, date, bullets in entries:
        draw.text((left, y), company, font=font(13), fill=orange)
        city_width = draw.textlength(city, font=font(10))
        draw.text((right - city_width, y + 1), city, font=font(10), fill=muted)
        y += 20
        draw.text((left, y), role, font=font(11, bold=True), fill=blue)
        date_width = draw.textlength(date, font=font(10))
        draw.text((right - date_width, y), date, font=font(10), fill=muted)
        y += 19
        for bullet in bullets:
            draw.ellipse((left + 2, y + 5, left + 6, y + 9), fill=blue)
            for line in wrap(bullet, right - left - 20, font(10)):
                draw.text((left + 16, y), line, font=font(10), fill=ink)
                y += 15
        y += 10

    y = heading("Ausbildung", y)
    for institution, degree, city, date in [
        (
            "Technische Universität München",
            "Master in Maschinenbau",
            "München",
            "2011 – 2013",
        ),
        (
            "RWTH Aachen University",
            "Bachelor in Verfahrenstechnik",
            "Aachen",
            "2007 – 2011",
        ),
    ]:
        draw.text((left, y), institution, font=font(12), fill=orange)
        width = draw.textlength(city, font=font(10))
        draw.text((right - width, y), city, font=font(10), fill=muted)
        y += 19
        draw.text((left, y), degree, font=font(11, bold=True), fill=blue)
        width = draw.textlength(date, font=font(10))
        draw.text((right - width, y), date, font=font(10), fill=muted)
        y += 27

    y = heading("Sprachen", y)
    draw.text((left, y), "Deutsch", font=font(11, bold=True), fill=blue)
    draw.text((left + 70, y), "Muttersprache   ●●●●●", font=font(10), fill=ink)
    draw.text((460, y), "Englisch", font=font(11, bold=True), fill=blue)
    draw.text((530, y), "B2   ●●●○○", font=font(10), fill=ink)

    draw.text((left, 1234), "lena-hoffmann.de", font=font(8), fill=muted)
    page = "Seite 1 / 1"
    draw.text(
        (right - draw.textlength(page, font=font(8)), 1234),
        page,
        font=font(8),
        fill=muted,
    )
    image.save(output_path, "PNG", optimize=True)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create Ivy League Lebenslauf Word templates",
    )
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--ats-output", type=Path)
    parser.add_argument("--preview-output", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    print(build_visual(args.output.resolve(), force=args.force))
    if args.ats_output:
        print(build_ats(args.ats_output.resolve(), force=args.force))
    if args.preview_output:
        print(build_preview(args.preview_output.resolve(), force=args.force))


if __name__ == "__main__":
    main()
